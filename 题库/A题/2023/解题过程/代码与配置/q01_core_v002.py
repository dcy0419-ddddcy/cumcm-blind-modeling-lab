"""Q1 approved prototype core v002. Offline, finite objects, no annual run entry."""
from __future__ import annotations
import math
from dataclasses import dataclass
from pathlib import Path
import numpy as np

EPS = np.finfo(float).eps
@dataclass(frozen=True)
class Tolerance:
    scale_m: float = 1000.0
    factor: float = 256.0
    @property
    def length(self): return self.factor * EPS * self.scale_m
    @property
    def direction(self): return self.factor * EPS

@dataclass
class Mirrors:
    centers: np.ndarray
    normals: np.ndarray
    u: np.ndarray
    v: np.ndarray
    widths: np.ndarray
    heights: np.ndarray
    vertical: np.ndarray
    @property
    def radii(self): return np.hypot(self.widths, self.heights) / 2
    def subset(self, ids):
        return Mirrors(*(getattr(self, f)[ids] for f in
                         ["centers","normals","u","v","widths","heights","vertical"]))

@dataclass(frozen=True)
class Cylinder:
    center: tuple = (0.,0.,80.)
    radius: float = 3.5
    height: float = 8.

class DomainError(ValueError): pass

def unit(x):
    a = np.asarray(x, dtype=float)
    norm = np.linalg.norm(a, axis=-1, keepdims=True)
    if np.any(norm == 0) or not np.all(np.isfinite(a)):
        raise DomainError("Zero or nonfinite direction")
    return a / norm

def solar(month, hour, latitude_deg=39.4, altitude_km=3.):
    lengths = [31,28,31,30,31,30,31,31,30,31,30,31]
    if not (1 <= month <= 12 and 0 <= hour <= 24):
        raise DomainError("Invalid representative date/time")
    D = sum(lengths[:month-1]) + 21 - 80
    delta = math.asin(math.sin(2*math.pi*D/365) * math.sin(math.radians(23.45)))
    w = math.pi / 12 * (hour - 12)
    phi = math.radians(latitude_deg)
    s = np.array([-math.cos(delta)*math.sin(w),
                  math.sin(delta)*math.cos(phi)-math.cos(delta)*math.cos(w)*math.sin(phi),
                  math.sin(delta)*math.sin(phi)+math.cos(delta)*math.cos(w)*math.cos(phi)])
    # Components implement the appendix with a signed east component; no acos branch.
    if abs(float(s@s)-1) > 1e-12 or s[2] <= 0:
        raise DomainError("Solar unit/positive-altitude domain failed")
    H = altitude_km
    a = .4237-.00821*(6-H)**2
    b = .5055+.00595*(6.5-H)**2
    c = .2711+.01858*(2.5-H)**2
    dni = 1.366*(a+b*math.exp(-c/s[2]))
    if not (math.isfinite(dni) and dni > 0):
        raise DomainError("Invalid DNI")
    return {"D":D,"delta":delta,"omega":w,"s0":s,"alpha":math.asin(s[2]),"dni":dni}

def edge_frame(normals):
    n = unit(normals)
    vertical = (n[:,0] == 0) & (n[:,1] == 0)
    u = np.column_stack((-n[:,1],n[:,0],np.zeros(len(n))))
    u[vertical] = [1,0,0]
    u = unit(u)
    v = np.cross(n,u)
    return u,v,vertical

def make_mirrors(centers, normals, width=6., height=6.):
    centers, normals = np.asarray(centers,float).reshape(-1,3), unit(np.asarray(normals,float).reshape(-1,3))
    u,v,vertical = edge_frame(normals)
    return Mirrors(centers,normals,u,v,
                   np.broadcast_to(width,(len(centers),)).copy(),
                   np.broadcast_to(height,(len(centers),)).copy(),vertical)

def aim_field(centers, s0, receiver):
    target = unit(np.asarray(receiver.center)-centers)
    return make_mirrors(centers, unit(target+s0)), target

def reflect(d,n):
    return d-2*np.sum(d*n,axis=-1,keepdims=True)*n

def cone_basis(s0):
    helper = np.eye(3)[int(np.argmin(np.abs(s0)))]
    e1 = unit(np.cross(s0,helper))
    return e1,np.cross(s0,e1)

def cone_directions(s0,beta,u,v):
    if not (0 <= beta < math.pi/2): raise DomainError("Cone reference plane domain")
    # 2 sin^2(beta/2) avoids loss in 1-cos(beta); no distribution clipping.
    mu = 1-np.asarray(u)*2*math.sin(beta/2)**2
    phi = 2*math.pi*np.asarray(v)
    e1,e2 = cone_basis(s0)
    radicand = (1-mu)*(1+mu)
    if np.any(radicand < 0): raise DomainError("Invalid cone radicand")
    transverse = np.sqrt(radicand)
    return mu[:,None]*s0 + transverse[:,None]*(np.cos(phi)[:,None]*e1+np.sin(phi)[:,None]*e2)

def check_domain(field, s0, beta, receiver, target):
    if not 0 <= beta < math.pi/2: raise DomainError("beta/reference plane")
    cosi = field.normals@s0
    if np.any(cosi > 1+64*EPS) or np.any(cosi < -1-64*EPS): raise DomainError("cosine range")
    # Avoid clipping: transverse norm computed directly rather than sqrt(1-cosi^2).
    transverse = np.linalg.norm(np.cross(field.normals,s0),axis=1)
    min_front = cosi*math.cos(beta)-transverse*math.sin(beta)
    min_horizon = s0[2]*math.cos(beta)-np.linalg.norm(s0[:2])*math.sin(beta)
    distances = np.linalg.norm(np.asarray(receiver.center)-field.centers,axis=1)
    start_margin = distances + math.hypot(receiver.radius,receiver.height/2)-field.radii
    if np.any(min_front <= 0) or min_horizon <= 0:
        raise DomainError("Whole-cone front/horizon domain failed; no truncation/renormalization")
    if np.any(distances > 1000) or np.any(start_margin <= 0):
        raise DomainError("Atmosphere or N04 starting-side domain failed")
    tau=.99321-.0001176*distances+1.97e-8*distances**2
    if np.any((tau<0)|(tau>1)): raise DomainError("Atmosphere outside physical bounds")
    return {"min_front":float(min_front.min()),"min_horizon":float(min_horizon),
            "min_end_start_margin_m":float(start_margin.min()),
            "min_outgoing_axis_dot":math.cos(beta),"strict_vertical_count":int(field.vertical.sum()),
            "min_horizontal_normal":float(np.linalg.norm(field.normals[:,:2],axis=1).min())}

def read_field(path):
    import openpyxl
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    ws=wb["Sheet1"]
    header=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
    if tuple(header)!=("x坐标 (m)","y坐标 (m)"): raise DomainError("Unexpected input header")
    rows=[]
    for row_no, values in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if len(values)!=2 or any(not isinstance(v,(float,int)) for v in values):
            raise DomainError(f"Bad row {row_no}")
        rows.append({"mirror_id":len(rows)+1,"excel_row":row_no,"x":float(values[0]),"y":float(values[1])})
    wb.close()
    centers=np.array([[r["x"],r["y"],4.] for r in rows])
    if len(rows)!=1745 or not np.all(np.isfinite(centers)): raise DomainError("Input count/nonfinite")
    return centers,rows

def conservative_candidates(field, i, axis, beta, tol):
    delta=field.centers-field.centers[i]
    dist=np.linalg.norm(delta,axis=1)
    axial=delta@axis
    perp=np.linalg.norm(delta-axial[:,None]*axis,axis=1)
    R=field.radii+field.radii[i]
    # If intersection exists: delta=q+lambda*d, |q|<=Ri+Rj,
    # lambda<=|delta|+R, and d in beta cone. Necessary conditions only.
    keep=(axial >= -R-tol.length)&(perp <= R+(dist+R)*math.sin(beta)+tol.length)
    keep[i]=False
    return np.flatnonzero(keep)

def rectangles(o,d,field,ids,limit,tol):
    """Nearest opaque rectangle; exact edges marked uncertain, not expanded."""
    q=len(o); limit=np.broadcast_to(limit,(q,))
    if not len(ids): return np.full(q,np.inf),np.zeros(q,bool),np.full(q,-1,int)
    f=field.subset(ids)
    den=np.einsum("qd,md->qm",d,f.normals)
    num=np.sum(f.centers*f.normals,axis=1)[None,:]-np.einsum("qd,md->qm",o,f.normals)
    t=np.full_like(den,np.inf)
    np.divide(num,den,out=t,where=den!=0)
    # Avoid inf*0 invalid values; parallel candidates are treated separately.
    tf=np.where(np.isfinite(t),t,0.)
    au=np.einsum("qd,md->qm",o,f.u)-np.sum(f.centers*f.u,axis=1)[None,:]+tf*np.einsum("qd,md->qm",d,f.u)
    av=np.einsum("qd,md->qm",o,f.v)-np.sum(f.centers*f.v,axis=1)[None,:]+tf*np.einsum("qd,md->qm",d,f.v)
    du=f.widths[None,:]/2-np.abs(au)
    dv=f.heights[None,:]/2-np.abs(av)
    forward=(t>tol.length)&(t<limit[:,None])
    inside=(du>=0)&(dv>=0)&(den!=0)
    hit=forward&inside
    near=(t>=-tol.length)&(t<=limit[:,None]+tol.length)&(du>=-tol.length)&(dv>=-tol.length)&(den!=0)
    uncertain=near&((du<=tol.length)|(dv<=tol.length)|(t<=tol.length)|
                    (np.isfinite(limit[:,None]) & (np.abs(tf-np.where(np.isfinite(limit),limit,0.)[:,None])<=tol.length))|(np.abs(den)<=tol.direction))
    # Coplanar rays: conservative bounding-ball overlap, never silently called a miss.
    delta=f.centers[None,:,:]-o[:,None,:]
    along=np.einsum("qmd,qd->qm",delta,d)
    perp2=np.sum(delta*delta,axis=2)-along*along
    possible_ball=(perp2<=f.radii[None,:]**2+tol.length**2)&(along+f.radii[None,:]>=0)&(along-f.radii[None,:]<=limit[:,None])
    uncertain |= (den==0)&(np.abs(num)<=tol.length)&possible_ball
    distances=np.where(hit,t,np.inf)
    j=np.argmin(distances,axis=1)
    nearest=distances[np.arange(q),j]
    which=np.where(np.isfinite(nearest),np.asarray(ids)[j],-1)
    return nearest,np.any(uncertain,axis=1),which

def cylinder_first(o,d,cyl,tol):
    """First finite-solid contact. Returns t, side/cap/exit/miss, uncertainty."""
    o=np.asarray(o,float); d=np.asarray(d,float); q=len(o)
    z=o-np.asarray(cyl.center)
    A=np.sum(d[:,:2]**2,axis=1); B=2*np.sum(z[:,:2]*d[:,:2],axis=1)
    C=np.sum(z[:,:2]**2,axis=1)-cyl.radius**2
    disc=B*B-4*A*C
    disc_tol=tol.factor*EPS*(B*B+4*np.abs(A*C)+cyl.radius**2)
    candidates=np.full((q,4),np.inf)
    near_candidates=np.full((q,4),np.inf)
    side_rad=np.zeros(q)
    good=(A>0)&(disc>=0)
    side_rad[good]=np.sqrt(disc[good])
    qq=-.5*(B+np.copysign(side_rad,B))
    t1=np.full(q,np.inf); t2=np.full(q,np.inf)
    np.divide(qq,A,out=t1,where=good)
    np.divide(C,qq,out=t2,where=good&(qq!=0))
    double=good&(qq==0); t2[double]=t1[double]
    for j,t in enumerate([t1,t2]):
        tf=np.where(np.isfinite(t),t,0.)
        zhit=z[:,2]+tf*d[:,2]
        inside=np.abs(zhit)<=cyl.height/2
        valid=good&(t>tol.length)&inside
        candidates[valid,j]=t[valid]
        near=good&(t>=-tol.length)&(np.abs(zhit)<=cyl.height/2+tol.length)&(
             (np.abs(np.abs(zhit)-cyl.height/2)<=tol.length)|
             (np.abs(disc)<=disc_tol)|(t<=tol.length))
        near_candidates[near,j]=np.maximum(t[near],0.)
    # Negative discriminant inside roundoff band: tangent-like uncertain contact.
    almost=(A>0)&(disc<0)&(disc>=-disc_tol)
    closest=np.zeros(q); np.divide(-B,2*A,out=closest,where=A>0)
    possible=almost&(closest>0)&(np.abs(z[:,2]+closest*d[:,2])<=cyl.height/2+tol.length)
    near_candidates[possible,0]=closest[possible]
    for j,capz in enumerate([-cyl.height/2,cyl.height/2],start=2):
        t=np.full(q,np.inf)
        np.divide(capz-z[:,2],d[:,2],out=t,where=d[:,2]!=0)
        tf=np.where(np.isfinite(t),t,0.)
        xy=z[:,:2]+tf[:,None]*d[:,:2]
        rr=np.linalg.norm(xy,axis=1)
        valid=(t>tol.length)&(rr<=cyl.radius)&(d[:,2]!=0)
        candidates[valid,j]=t[valid]
        near=(d[:,2]!=0)&(t>=-tol.length)&(rr<=cyl.radius+tol.length)&(
              (np.abs(rr-cyl.radius)<=tol.length)|(t<=tol.length)|(np.abs(d[:,2])<=tol.direction))
        near_candidates[near,j]=np.maximum(t[near],0.)
    j=np.argmin(candidates,axis=1); first=candidates[np.arange(q),j]
    touch=np.min(near_candidates,axis=1)
    uncertain=np.isfinite(touch)&(touch<=first+tol.length)
    kind=np.where(np.isfinite(first),np.where(j<2,"side","cap"),"miss")
    inside_origin=(C<0)&(np.abs(z[:,2])<cyl.height/2)
    kind=np.where(inside_origin&np.isfinite(first),"exit",kind)
    # Exterior-entry direction for active side; tangent cannot count as entry.
    tf=np.where(np.isfinite(first),first,0.)
    xy=z[:,:2]+tf[:,None]*d[:,:2]
    radial_dot=np.sum(xy*d[:,:2],axis=1)/cyl.radius
    active=(kind=="side")&(radial_dot < -tol.direction)
    return first,kind,active,uncertain

def trace(o,s,field,i,s0,target,cyl,beta,tol,screen=True,end_offset=0.):
    r=reflect(-s,field.normals[i])
    axis=target[i]
    denom=r@axis
    Rb=math.hypot(cyl.radius,cyl.height/2)+end_offset
    if end_offset<0: raise DomainError("Test endpoint may only extend approved envelope")
    numer=Rb-(o-np.asarray(cyl.center))@axis
    if np.any(denom<=0) or np.any(numer<=0): raise DomainError("N04 no positive endpoint")
    end=numer/denom
    all_ids=np.delete(np.arange(len(field.centers)),i)
    inc_ids=conservative_candidates(field,i,s0,beta,tol) if screen else all_ids
    out_ids=conservative_candidates(field,i,axis,beta,tol) if screen else all_ids
    it,iu,ii=rectangles(o,s,field,inc_ids,np.inf,tol)
    ic,ik,ia,icu=cylinder_first(o,s,cyl,tol)
    S=~(np.isfinite(it)|np.isfinite(ic))
    rt,rk,active,rcu=cylinder_first(o,r,cyl,tol)
    limit=np.minimum(rt,end)
    bt,bu,bi=rectangles(o,r,field,out_ids,limit,tol)
    B=~np.isfinite(bt)
    R=active
    unknown=iu|icu|(S&(bu|rcu))
    return {"S":S,"B":B,"R":R,"unknown":unknown,"incoming_object":ii,
            "blocking_object":bi,"receiver_kind":rk,
            "candidate_in":len(inc_ids),"candidate_out":len(out_ids)}

def random_rays(field,i,s0,beta,n,seed_words):
    rng=np.random.Generator(np.random.PCG64(np.random.SeedSequence(seed_words)))
    a=rng.random((n,4))
    p=field.centers[i]+((a[:,0]-.5)*field.widths[i])[:,None]*field.u[i]+((a[:,1]-.5)*field.heights[i])[:,None]*field.v[i]
    s=cone_directions(s0,beta,a[:,2],a[:,3])
    return p,s

def quadrature_rays(field,i,s0,beta,na,nmu,nphi,phase=0.):
    a=(np.arange(na)+.5)/na-.5
    mu=(np.arange(nmu)+.5)/nmu
    ph=((np.arange(nphi)+.5+phase)/nphi)%1
    x,y,u,v=np.meshgrid(a,a,mu,ph,indexing="ij")
    p=field.centers[i]+(x.ravel()*field.widths[i])[:,None]*field.u[i]+(y.ravel()*field.heights[i])[:,None]*field.v[i]
    s=cone_directions(s0,beta,u.ravel(),v.ravel())
    return p,s # equal product weights in actual area and mu/phi measure

def summarize(g,S,B,R,unknown,cosi,area,dni,tau,rho=.92,certificate=None):
    g=np.asarray(g,float); unknown=np.asarray(unknown,bool)
    if len(g)==0 or not np.all(np.isfinite(g)) or np.any(g<0):
        raise DomainError("Invalid weights")
    if not all(math.isfinite(x) for x in [cosi,area,dni,tau,rho]) or area<=0 or dni<0:
        raise DomainError("Invalid energy inputs")
    f1=np.asarray(S)&np.asarray(B); f2=f1&np.asarray(R)
    X=np.column_stack((g,g*f1,g*f2))
    sums=X.sum(axis=0); raw=X.mean(axis=0)
    scale=rho*dni*area
    pi0=scale*cosi
    result={"samples":len(g),"raw_sums":sums.tolist(),"raw_cross_sums":(X.T@X).tolist(),
            "raw_mean":raw.tolist(),"raw_energy":(scale*raw).tolist(),
            "cosine":float(cosi),"tau":float(tau),"rho":float(rho),"area_m2":float(area),
            "dni_kw_m2":float(dni),"analytic_pi0_kw":float(pi0),
            "raw_projection_residual":float(raw[0]-cosi),
            "counts":{"shadow":int((~np.asarray(S)).sum()),"survive":int(f1.sum()),
                      "capture":int(f2.sum()),"unknown":int(unknown.sum())}}
    if sums[0] <= 0 or pi0==0:
        result.update(status="TRUE_ZERO_INPUT" if pi0==0 else "INPUT_FAILURE",
                      eta_sb=None,eta_trunc=None,eta_total=None,power_kw=0. if pi0==0 else None)
        return result
    f1hat,f2hat=sums[1]/sums[0],sums[2]/sums[0]
    eta_total=rho*cosi*tau*f2hat
    status="ESTIMATED"
    eta_sb=float(f1hat); eta_trunc=float(sums[2]/sums[1]) if sums[1]>0 else None
    power=tau*pi0*f2hat
    if sums[1]==0:
        status="TRUE_ZERO_SURVIVOR" if certificate=="zero_survivor" else "SAMPLED_ZERO_SURVIVOR"
        if certificate!="zero_survivor": eta_sb=None; eta_total=None; power=None
    elif sums[2]==0:
        status="TRUE_ZERO_CAPTURE" if certificate=="zero_capture" else "SAMPLED_ZERO_CAPTURE"
        if certificate!="zero_capture": eta_trunc=None; eta_total=None; power=None
    if np.any(unknown):
        status="BOUNDARY_UNCERTAIN"
        eta_sb=eta_trunc=eta_total=power=None
    # A point estimator is retained under its numerical name even when not certifiable.
    low2=float(np.sum(g*f2*(~unknown))/sums[0])
    up2=low2+float(np.sum(g*unknown)/sums[0])
    result.update(status=status,eta_sb=eta_sb,eta_trunc=eta_trunc,eta_total=eta_total,power_kw=power,
                  estimate_f1=float(f1hat),estimate_f2=float(f2hat),
                  estimate_power_kw=float(tau*pi0*f2hat),
                  estimated_pi1_kw=float(pi0*f1hat),estimated_pi2_kw=float(pi0*f2hat),
                  boundary_capture_fraction_interval=[low2,up2],
                  boundary_weight_fraction=float(np.sum(g*unknown)/sums[0]))
    # Delta-method diagnostic SE, not a confidence guarantee for rare/zero events.
    if len(g)>1:
        cov=np.cov(X,rowvar=False,ddof=1)/len(g)
        grad=np.array([-raw[2]/raw[0]**2,0,1/raw[0]])
        v=float(grad@cov@grad)
        if v < -1e-15: raise DomainError("Negative variance beyond roundoff")
        result["delta_se_f2"]=math.sqrt(max(v,0.))
        if raw[1]>0:
            gt=np.array([0,-raw[2]/raw[1]**2,1/raw[1]])
            vv=float(gt@cov@gt)
            if vv < -1e-15: raise DomainError("Negative ratio variance")
            result["delta_se_trunc"]=math.sqrt(max(vv,0.))
    return result

def evaluate(field,i,s0,target,cyl,beta,tol,o,s,dni=1.,tau=1.,screen=True,certificate=None):
    if np.any(s@s0<=0) or np.any(s@field.normals[i]<=0):
        raise DomainError("Ray weight front/reference domain")
    g=(s@field.normals[i])/(s@s0)
    ev=trace(o,s,field,i,s0,target,cyl,beta,tol,screen)
    ans=summarize(g,ev["S"],ev["B"],ev["R"],ev["unknown"],float(field.normals[i]@s0),
                  field.widths[i]*field.heights[i],dni,tau,certificate=certificate)
    ans["candidate_counts"]=[ev["candidate_in"],ev["candidate_out"]]
    abnormal=np.flatnonzero(ev["unknown"])[:8]
    ans["abnormal_examples"]=[{"sample":int(j),"point":o[j].tolist(),"sun_direction":s[j].tolist(),
                              "receiver_kind":str(ev["receiver_kind"][j])} for j in abnormal]
    return ans,ev

def fixed_mean(values):
    """None/NaN stays undefined on fixed original denominator."""
    if not values or any(v is None or not math.isfinite(v) for v in values): return None
    return float(sum(values)/len(values))

def h02_summary(eta_components,dni,areas,direct_total=None):
    """Generic complete month x time x mirror array. Used only on synthetic data here."""
    arr=np.asarray(eta_components,float)
    if arr.ndim!=4 or arr.shape[-1]!=5: raise DomainError("Expected month,time,mirror,five factors")
    dni=np.asarray(dni,float); areas=np.asarray(areas,float)
    if dni.shape!=arr.shape[:2] or areas.shape!=(arr.shape[2],): raise DomainError("H02 shape")
    if np.any(areas<=0): raise DomainError("H02 area")
    # np.mean propagates NaN; no nanmean or object deletion.
    total=np.prod(arr,axis=-1)
    if direct_total is not None:
        direct=np.asarray(direct_total,float)
        if direct.shape!=total.shape: raise DomainError("Direct total shape")
        valid=np.isfinite(total)
        if np.any(np.abs(direct[valid]-total[valid])>1e-12):
            raise DomainError("Direct total contradicts defined factor product")
        total=np.where(valid,total,direct)
    power=dni*np.sum(total*areas,axis=-1)
    return {"component_month":np.mean(arr,axis=(1,2)),"total_month":np.mean(total,axis=(1,2)),
            "component_year":np.mean(arr,axis=(0,1,2)),"total_year":float(np.mean(total)),
            "power_month":np.mean(power,axis=1),"power_year":float(np.mean(power)),
            "unit_power_year":float(np.mean(power)/areas.sum())}
