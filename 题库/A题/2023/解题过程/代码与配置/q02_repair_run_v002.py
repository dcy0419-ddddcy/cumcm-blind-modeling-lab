"""Explicit bounded repair actions; no optical scoring on import before freeze."""
import time
START=time.perf_counter()
from pathlib import Path
import sys,json,math,traceback,copy
HERE=Path(__file__).resolve().parent;sys.path.insert(0,str(HERE))
import q02_repair_common_v001 as io
action=sys.argv[1] if len(sys.argv)>1 else 'invalid'
budget=io.Budget('repair_'+action,'confirmation' if action=='confirm' else 'search',START)
try:
 import numpy as np
 import q02_repair_engine_v001 as eng
 O=io.OUT;C=O/'candidates';F=eng.f;D=F.design_module;CP=O/'搜索与确认冻结配置-v002.json'
 if action=='init':
  if CP.exists():raise RuntimeError('CONFIG_EXISTS')
  cfg={'version':'repair-v001','created':io.now(),'root_seed':2026090505,'before_first_new_optical_score':True,'parent_failure':{'candidate':'C22','path':'../Q02-搜索-v001/独立确认结论.json','sha256':io.sha(O.parent/'Q02-搜索-v001/独立确认结论.json'),'used_as_search_data':True},'search_namespaces':[1900,1910],
   'low':{'times':eng.TIMES,'B':2,'n':16,'levels':[8,16],'namespace':1900},
   'fine':{'times':eng.TIMES,'B':4,'n':64,'levels':[32,64],'namespace':1910},
   'confirmation':{'B':8,'n':256,'levels':[128,256],'namespace':1990,'fixed_sample_no_adaptive_stop':True,'min_reserve_seconds':1200,'U':'max(t7 whole-batch jackknife halfwidth, final-half-prefix difference, final-raw-integral difference) + coordinatewise full-range low-survivor(<100) influence bound','efficiency_abs':.001,'power_relative':.005,'rating':'annual P-U_P>=60000kW; all78 report metrics and no required unresolved states','low_survival_threshold':100,'zero_rule':'no zero/one fill, keep fixed H02, unknown propagates, no seed retries','interval':'approximate pointwise diagnostic only, no simultaneous or physical guarantee'},
   'rules':{'candidate_names':'R023 onward, generator signature excludes name/version; each spec saved before scoring','family':'six-sector straight-edge bands with proven triangular lattice spacing; full fixed-field/tower exclusion clip and all-pair check','initial':[{'w':6.75,'h':6.75,'z':6,'tower':[0,0],'phi':math.pi/6,'margin':.05},{'w':7,'h':7,'z':6,'tower':[0,0],'phi':math.pi/6,'margin':.05}],
    'allowed_moves':{'width_height':'2..8 and w>=h; local +/- .05,.1,.15,.25,.5 m, joint dimensions or height-only; z6 retained first, optionalz5.5 if h<11','tower':'from parent +/-(0,15),(0,30),(15,0); |x|<=30, y in[-60,30]','phi':'0,pi/12,pi/6','spacing_margin':'.01 or .05 m; numerical search restriction, not engineering clearance'},
    'selection':'Until fine-gated candidate exists rank valid whole60 point P first; also retain q/area. Once exists retain baseline and rank q among fine-gated. At least one q-directed local adjustment or complementary full60 comparison; do not start local only when feasible.',
    'NA_at_low':'low total formal point may be NA due pooled zero; record heuristic score separately for shortlist only, no rating. At fine unresolved may trigger one higher sample stage only if budget and new registered config; otherwise not eligible.',
    'shortlist':'promising low P and q receive fine; cannot call noise-scale q differences established improvements; paired batch comparison diagnostic retained','fine_gate':'all required point quantities defined; annual P - max(50kW, 2*max(3.182446305*JK_SE,abs(P-prefixP),abs(P-rawP))) >60000. Search risk margin only, not rated rule or strict confidence. 50kW floor exceeds previous between-level ~10kW and U~6kW; intentionally conservative management choice.',
    'budgets':{'maximum_new_candidates':8,'maximum_low':6,'maximum_fine':3,'search_seconds':2400,'total_seconds':3600,'reserve_confirmation':1200,'startup_low_seconds':310,'startup_fine_seconds':380,'cost_rule':'last comparable elapsed times1.12 times N ratio; reserve20sec freeze; no new stage if prediction exceeds phase cap; projections not hard guarantees'},
    'rounding':'x/y10 decimal before scoring and Excel; exact readback frozen','CRN':'PCG64 SeedSequence(root,namespace,time index,stable position_key,batch). Common keys receive same normalized source points/directions, not guaranteed same physicalmirror; ordering does not drive RNG. low andfine independent, halfprefix correlated and countedonce.'},
   'bindings':{p.name:io.sha(p)for p in [Path(__file__),HERE/'q02_repair_engine_v001.py',HERE/'q02_repair_common_v001.py',HERE/'q02_fast_v002.py',HERE/'q02_compact_v001.py',HERE/'q02_search_summary_v002.py',HERE/'q02_hex_bands_v001.py']}}
  io.save(CP,cfg);io.save(O/'自适应队列.json',{'next_number':23,'candidates':[],'selection_log':[]})
  # Source-equivalence audit for numerical execution: output binding/guards only.
  old=(HERE/'q02_search_engine_v002.py').read_text(encoding='utf-8');new=(HERE/'q02_repair_engine_v001.py').read_text(encoding='utf-8')
  restored=new.replace('import q02_repair_common_v001 as io','import q02_search_common_v001 as io').replace('  budget.guard(12)\n','').replace(' budget.guard(15)\n','')
  if restored!=old:raise RuntimeError('ENGINE_DIFF_OUTSIDE_EXPECTED_SCOPE')
  io.save(O/'最小改造静态核对.json',{'exact_old_source_restored':True,'changes':['output directory/budget binding','budget guards before chunk reconstruction and stacking'],'physics_statistics_unchanged':True,'no_new_numeric_kernel':True})
 elif action in ['rate','freeze','confirm']:
  cfg=io.load(CP)
  for nm,dg in cfg['bindings'].items():
   if io.sha(HERE/nm)!=dg:raise RuntimeError('FROZEN_SOURCE_CHANGED: '+nm)
  if action=='rate':
   if (O/'最终拟提交候选冻结.json').exists():raise RuntimeError('NO_SEARCH_AFTER_FREEZE')
   kind=sys.argv[2]
   if kind not in ['low','fine']:raise ValueError('unknown fidelity')
   request=json.loads(sys.argv[3]);queue=io.load(O/'自适应队列.json')
   if 'name' in request:
    item=next(v for v in queue['candidates']if v['name']==request['name']);spec=item['spec']
   else:
    # Only approved frozen move range; explicit full signature prevents duplicate re-labelling.
    w,h,z=request['w'],request['h'],request.get('z',6);tx,ty=request.get('tower',[0,0]);phi=request.get('phi',math.pi/6);margin=request.get('margin',.05)
    if not(2<=h<=w<=8 and z in [5.5,6]and z>h/2 and abs(tx)<=30 and -60<=ty<=30 and phi in [0,math.pi/12,math.pi/6]and margin in [.01,.05]):raise ValueError('OUTSIDE_FROZEN_SEARCH_RANGE')
    sig=[w,h,z,tx,ty,phi,margin]
    if any(v['signature']==sig for v in queue['candidates']):raise RuntimeError('DUPLICATE_PARAMETERS')
    if len(queue['candidates'])>=cfg['rules']['budgets']['maximum_new_candidates']:raise RuntimeError('CANDIDATE_LIMIT')
    name='R%03d'%queue['next_number'];queue['next_number']+=1
    spec=eng.hex_bands.make_hex_spec((tx,ty),w,h,z,phi=phi,spacing_margin=margin,name=name)
    item={'name':name,'signature':sig,'spec':spec,'parent':request.get('parent','C22'),'reason':request['reason'],'registered_before_output':io.now()};queue['candidates'].append(item);io.save(O/'自适应队列.json',queue)
   existing=list(C.glob('*/'+kind+'/score.json'))
   if not(C/spec['name']/kind/'score.json').exists()and len(existing)>=cfg['rules']['budgets']['maximum_'+kind]:raise RuntimeError('FIDELITY_COUNT_LIMIT')
   start=time.perf_counter();folder=C/spec['name'];fd=eng.generate(spec,folder,budget)
   # Each changed design checks dense same-ray event equivalence at one preselected time.
   if not(folder/'困难组合同射线.json').exists():
    sc=F.scene_at(fd,12,9);pairs=[F.candidate_pair(sc,j)for j in range(fd.n)];idx=max(range(fd.n),key=lambda j:(len(pairs[j].incoming)+len(pairs[j].outgoing),-fd.mirror_ids[j]))
    io.save(folder/'困难组合预登记.json',{'month':12,'hour':9,'index':idx,'mirror_id':fd.mirror_ids[idx],'rule':'maximum sum conservative candidate counts, before rays','candidates':[len(pairs[idx].incoming),len(pairs[idx].outgoing)]})
    origins,dirs=F.c.random_rays(sc.mirrors,idx,sc.s0,sc.beta,64,[2026090505,1881,12]);g,ev=F.trace_same(sc,idx,origins,dirs);g2,ev2=F.trace_same(sc,idx,origins,dirs,screen=False)
    checks={k:bool(np.array_equal(ev[k],ev2[k]))for k in F.EVENT_KEYS};checks['g']=bool(np.array_equal(g,g2));io.save(folder/'困难组合同射线.json',{'checks':checks,'all_pass':all(checks.values()),'source_samples':64})
    if not all(checks.values()):raise RuntimeError('SCREENED_EXHAUSTIVE_MISMATCH')
   plan=cfg[kind];estimate=cfg['rules']['budgets']['startup_'+kind+'_seconds']
   if existing:
    prior=io.load(existing[-1]);estimate=prior['evaluation_seconds']*1.12*fd.n/prior['n']
   if budget.used()+estimate+20>=2400:raise RuntimeError('STARTUP_ESTIMATE_EXCEEDS_SEARCH_RESERVE')
   arrays,recs=eng.evaluate(fd,folder,kind,eng.TIMES,plan['B'],plan['n'],plan['levels'],plan['namespace'],budget)
   budget.guard(20);summary=eng.stored_summary(arrays,folder/kind,'full_comparison');score=eng.quick_score(arrays)
   point=summary['point'][-1];se=summary['jackknife_se'][-1];prefix=summary['prefix_half_point'][-1];raw=summary['raw_integral_alternative'][-1]
   required_na=[{'row':i,'metric':j} for i,row in enumerate(summary['point']) for j,value in enumerate(row) if not np.isfinite(float(value) if value is not None else float('nan'))]
   defined=np.isfinite(np.asarray(point,dtype=float)).all();margin=None;gate=False
   if defined and np.isfinite([se[4],prefix[4],raw[4]]).all():
    margin=max(50.,2*max(3.182446305*float(se[4]),abs(point[4]-prefix[4]),abs(point[4]-raw[4])));gate=bool(kind=='fine' and point[4]-margin>60000 and not required_na)
   score.update(name=fd.name,design_key=fd.key,n=fd.n,area=fd.total_area,spec=spec,score_kind=kind,point=point,jackknife_se=se,states=summary['pooled_state_counts'],required_na=required_na,fine_gate_margin_kw=margin,fine_gate=gate,evaluation_seconds=time.perf_counter()-start,actual_full60=True)
   io.save(folder/kind/'score.json',score);io.event('搜索轨迹.jsonl',score);budget.tick();print(json.dumps(io.convert(score),ensure_ascii=False))
  elif action=='freeze':
   name=sys.argv[2];best=io.load(C/name/'fine/score.json')
   if not best['fine_gate']:raise RuntimeError('NO_POSITIVE_MARGIN_FINE_EVIDENCE')
   if (O/'最终拟提交候选冻结.json').exists():raise RuntimeError('ALREADY_FROZEN')
   budget.guard(20)
   from openpyxl import Workbook,load_workbook
   fd,_=eng.load_design(C/name);des=fd.to_design();dest=O/'frozen';dest.mkdir(exist_ok=True)
   wb=Workbook();ws=wb.active;ws.title='设计参数';ws.append(['参数','值','单位'])
   pars={'tower_x':des.tower_xy[0],'tower_y':des.tower_xy[1],'width':des.width,'height':des.height,'installation_height':des.installation_height,'n':des.n}
   for k,v in pars.items():ws.append([k,v,'count'if k=='n'else'm'])
   ms=wb.create_sheet('逐镜设计');ms.append(['序号','x坐标 (m)','y坐标 (m)','安装高度 (m)','镜面宽 (m)','镜面高 (m)'])
   tmpl=load_workbook(io.ROOT/'附件/附件01.xlsx',read_only=True,data_only=True);header=list(next(tmpl.active.values));tmpl.close();os=wb.create_sheet('镜场设计');os.append(header)
   for i,r in enumerate(des.mirrors,1):
    ms.append([i,r['x'],r['y'],des.installation_height,des.width,des.height]);os.append([*des.tower_xy,i,des.width,des.height,r['x'],r['y'],des.installation_height])
   wb.active=2;path=dest/(name+'-待确认设计.xlsx');wb.save(path)
   rb=load_workbook(path,data_only=True,read_only=True);pp={r[0]:r[1]for r in list(rb['设计参数'].values)[1:]};mr=list(rb['逐镜设计'].values)[1:];rr=list(rb['镜场设计'].values)[1:];rb.close()
   if pp!=pars or len(mr)!=des.n or len(rr)!=des.n:raise RuntimeError('EXCEL_PARAMETERS_COUNT_FAILURE')
   mapped=[]
   for i,row in enumerate(mr):
    if row!=(i+1,des.mirrors[i]['x'],des.mirrors[i]['y'],des.installation_height,des.width,des.height)or rr[i]!=(*des.tower_xy,i+1,des.width,des.height,row[1],row[2],des.installation_height):raise RuntimeError('EXACT_EXPORT_READBACK_FAILED')
    mapped.append({'mirror_id':des.mirrors[i]['mirror_id'],'position_key':des.mirrors[i]['position_key'],'x':row[1],'y':row[2],'export_row':i+2})
   actualdes=D.Design(des.name,'confirmed-input-v001',des.tower_xy,des.width,des.height,des.installation_height,mapped,des.metadata);actual=F.freeze_design(actualdes);eng.compact.save_compact(actualdes,dest)
   io.save(dest/'身份映射.json',[{'export_index':i+1,'mirror_id':r['mirror_id'],'position_key':r['position_key']}for i,r in enumerate(mapped)])
   io.save(dest/'交付精度读回.json',{'excel':str(path),'sha256':io.sha(path),'n':len(mapped),'actual_design_key':actual.key,'geometry':actual.validation,'max_coordinate_difference_m':float(np.max(np.abs(actual.centers-fd.centers))),'nominal_original_key':fd.key,'actual_total_area':actual.total_area})
   io.save(O/'最终拟提交候选冻结.json',{'config_file':CP.name,'name':name,'created':io.now(),'before_confirmation':True,'estimated_rating_met_before_confirmation':True,'design_key':actual.key,'bundle_sha256':io.sha(dest/'design.bundle.json'),'excel_file':'frozen/'+path.name,'excel_sha256':io.sha(path),'best_comparison':best,'confirmation_config':cfg['confirmation'],'search_budget_used_seconds':budget.used(),'confirmation_data_used_for_selection':False,'ranking_basis':'preserve refined positive-margin candidate; compare q only at fine fidelity; finite search not optimum'})
   print('FROZEN '+name)
  else:
   fr=io.load(O/'最终拟提交候选冻结.json');plan=fr['confirmation_config']
   if plan!=cfg['confirmation']or plan['namespace']in cfg['search_namespaces']:raise RuntimeError('CONFIRMATION_BINDING')
   if io.sha(O/fr['excel_file'])!=fr['excel_sha256']:raise RuntimeError('EXCEL_CHANGED')
   fd,_=eng.load_design(O/'frozen')
   if fd.key!=fr['design_key']or io.sha(O/'frozen/design.bundle.json')!=fr['bundle_sha256']:raise RuntimeError('DESIGN_CHANGED')
   arrays,recs=eng.evaluate(fd,O/'frozen','confirmation',eng.TIMES,plan['B'],plan['n'],plan['levels'],plan['namespace'],budget)
   budget.guard(30);ans=eng.stored_summary(arrays,O/'frozen/confirmation','full_confirmation')
   io.save(O/'独立确认结论.json',{'name':fd.name,'design_key':fd.key,'decision':ans['formal_decision'],'annual':ans['point'][-1],'annual_U':ans['conservative_work_indicator'][-1],'all_precision':ans['all_table_items_precision_met'],'rating_lower_kw':ans['rated_power_lower_work_value_kw'],'unresolved':ans['required_table_na_locations'],'low_survival':ans['low_survival'],'numeric_checks':ans['numeric_checks'],'reconstructed_from_saved_chunks':True,'independent_batch_count':plan['B'],'full_combinations':fd.n*60,'source_samples':fd.n*60*plan['B']*plan['n']})
   print(json.dumps(io.load(O/'独立确认结论.json'),ensure_ascii=False))
 else:raise ValueError('init/rate/freeze/confirm required')
 budget.finish()
except BaseException:
 io.event('失败运行.jsonl',{'time':io.now(),'action':action,'error':traceback.format_exc(),'used':budget.used()});budget.finish('FAILED');raise
