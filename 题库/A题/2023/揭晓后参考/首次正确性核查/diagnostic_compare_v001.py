from __future__ import annotations

import hashlib
import json
import math
import time
import csv
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

W = Path(r"C:\Users\admin\OneDrive\Desktop\26国赛-单题盲解区\A-P8C3")
P = W / "工作记录" / "揭晓后对照" / "首次正确性核查-v001"
SOURCE = W / "工作记录" / "论文" / "全题整合" / "结果与来源-v001.json"
REFS = P / "参考论文" / "参考表格人工转录-v001.json"
OUT = P / "诊断结果-v001.json"
MONTHLY_CSV = P / "Q1逐月差值-v001.csv"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def diff(local: float, ref: float) -> dict:
    absolute = local - ref
    return {
        "local": local,
        "reference": ref,
        "local_minus_reference": absolute,
        "relative_to_reference": None if ref == 0 else absolute / ref,
    }


def read_design(path: Path, kind: str) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[2] if kind == "q2" else wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if kind == "q2":
        # tower_x, tower_y, id, width, height, x, y, z
        tower = np.array([float(rows[0][0]), float(rows[0][1])])
        ids = np.array([int(r[2]) for r in rows])
        widths = np.array([float(r[3]) for r in rows])
        heights = np.array([float(r[4]) for r in rows])
        xy = np.array([[float(r[5]), float(r[6])] for r in rows])
        z = np.array([float(r[7]) for r in rows])
    else:
        tower = np.array([float(rows[0][0]), float(rows[0][1])])
        ids = np.array([int(r[2]) for r in rows])
        widths = np.array([float(r[3]) for r in rows])
        heights = np.array([float(r[4]) for r in rows])
        xy = np.array([[float(r[5]), float(r[6])] for r in rows])
        z = np.array([float(r[7]) for r in rows])

    n = len(rows)
    min_d = math.inf
    min_pair = None
    min_margin = math.inf
    min_margin_pair = None
    for i0 in range(0, n, 256):
        i1 = min(n, i0 + 256)
        dxy = xy[i0:i1, None, :] - xy[None, :, :]
        ds = np.sqrt(np.sum(dxy * dxy, axis=2))
        for ii in range(i1 - i0):
            ds[ii, : i0 + ii + 1] = np.inf
        flat = int(np.argmin(ds))
        ii, jj = np.unravel_index(flat, ds.shape)
        if ds[ii, jj] < min_d:
            min_d = float(ds[ii, jj])
            min_pair = [int(ids[i0 + ii]), int(ids[jj])]
        thresholds = np.maximum(widths[i0:i1, None], widths[None, :]) + 5.0
        margins = ds - thresholds
        flatm = int(np.argmin(margins))
        im, jm = np.unravel_index(flatm, margins.shape)
        if margins[im, jm] < min_margin:
            min_margin = float(margins[im, jm])
            min_margin_pair = [int(ids[i0 + im]), int(ids[jm])]

    radii = np.linalg.norm(xy, axis=1)
    tower_d = np.linalg.norm(xy - tower[None, :], axis=1)
    clearance = z - heights / 2.0
    return {
        "path": str(path.relative_to(W)).replace("\\", "/"),
        "sha256": sha256(path),
        "N": n,
        "ids_unique": len(set(ids.tolist())) == n,
        "tower_xy_m": tower.tolist(),
        "total_area_m2": math.fsum((widths * heights).tolist()),
        "max_center_radius_m": float(np.max(radii)),
        "field_center_margin_m": float(350.0 - np.max(radii)),
        "min_tower_center_distance_m": float(np.min(tower_d)),
        "tower_exclusion_margin_m": float(np.min(tower_d) - 100.0),
        "min_pair_distance_m": min_d,
        "min_pair_ids": min_pair,
        "min_spacing_margin_m": min_margin,
        "min_spacing_margin_pair_ids": min_margin_pair,
        "min_ground_clearance_m": float(np.min(clearance)),
        "all_finite": bool(np.isfinite(xy).all() and np.isfinite(widths).all() and np.isfinite(heights).all() and np.isfinite(z).all()),
        "all_constraints_pass": bool(
            np.max(radii) <= 350.0
            and np.min(tower_d) >= 100.0
            and min_margin >= -1e-12
            and np.min(clearance) > 0.0
        ),
        "width_values_m": sorted(set(widths.tolist())),
        "height_values_m": sorted(set(heights.tolist())),
        "installation_height_values_m": sorted(set(z.tolist())),
    }


def main() -> None:
    started = time.perf_counter()
    local = json.loads(SOURCE.read_text(encoding="utf-8"))
    refs = json.loads(REFS.read_text(encoding="utf-8"))
    q1 = local["questions"]["q1"]
    labels = ["optical", "cosine", "shadow_block", "truncation", "power_mw", "unit_power"]
    q1_annual = {}
    q1_monthly = {}
    for rid in ("A0165", "A0127", "A092"):
        local_annual = list(map(float, q1["annual"]["point"]))
        local_annual[4] /= 1000.0
        q1_annual[rid] = {
            label: diff(local_annual[i], float(refs[rid]["q1_annual"][i]))
            for i, label in enumerate(labels)
        }
        q1_monthly[rid] = []
        for month in range(12):
            q1_monthly[rid].append({
                "month": month + 1,
                **{
                    label: diff(float(q1["point"][month][i]), float(refs[rid]["q1_monthly"][month][i]))
                    for i, label in enumerate(["optical", "cosine", "shadow_block", "truncation", "unit_power"])
                },
            })

    q2 = local["questions"]["q2"]
    q3 = local["questions"]["q3"]
    design_q2 = read_design(W / "工作记录/诊断结果/Q02-修复-v001/result2.xlsx", "q2")
    design_q3 = read_design(W / "工作记录/诊断结果/Q03-实施-v001/delivery/result3.xlsx", "q3")
    comparisons = {}
    for rid in ("A0165", "A0127", "A092"):
        comparisons[rid] = {
            "q2_power_mw": diff(q2["annual"]["power_mw"], refs[rid]["q2"]["power_mw"]),
            "q2_unit_power": diff(q2["annual"]["q_kw_m2"], refs[rid]["q2"]["q_kw_m2"]),
            "q3_power_mw": diff(q3["annual"]["power_mw"], refs[rid]["q3"]["power_mw"]),
            "q3_unit_power": diff(q3["annual"]["q_kw_m2"], refs[rid]["q3"]["q_kw_m2"]),
        }

    elapsed = time.perf_counter() - started
    out = {
        "scope": "No optical rerun. Saved-statistics arithmetic and independent formal-workbook geometry/readback only.",
        "inputs": {"local": {"path": str(SOURCE.relative_to(W)).replace("\\", "/"), "sha256": sha256(SOURCE)}, "reference_transcription": {"path": str(REFS.relative_to(W)).replace("\\", "/"), "sha256": sha256(REFS)}},
        "q1_annual_differences": q1_annual,
        "q1_monthly_differences": q1_monthly,
        "q2_q3_reported_value_differences": comparisons,
        "formal_design_readback": {"q2": design_q2, "q3": design_q3},
        "algebra": {
            "q1_P_minus_Aq_kw": q1["annual"]["power_kw"] - q1["total_area_m2"] * q1["annual"]["q_kw_m2"],
            "q2_P_minus_Aq_kw": q2["annual"]["power_kw"] - q2["total_area_m2"] * q2["annual"]["q_kw_m2"],
            "q3_P_minus_Aq_kw": q3["annual"]["power_kw"] - q3["total_area_m2"] * q3["annual"]["q_kw_m2"],
            "q3_rating_margin_kw": q3["annual"]["rating_margin_kw"],
            "paired_relative_q_improvement": local["comparison"]["relative_q_difference"],
            "paired_difference_minus_indicator_kw_m2": local["comparison"]["difference_minus_work_indicator"],
        },
        "elapsed_seconds": elapsed,
        "new_optical_rays": 0,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    with MONTHLY_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["reference_id", "month", "metric", "local", "reference", "local_minus_reference", "relative_to_reference"])
        for rid, monthly in q1_monthly.items():
            for row in monthly:
                for metric in ("optical", "cosine", "shadow_block", "truncation", "unit_power"):
                    d = row[metric]
                    writer.writerow([rid, row["month"], metric, d["local"], d["reference"], d["local_minus_reference"], d["relative_to_reference"]])
    print(json.dumps({"output": str(OUT), "elapsed_seconds": elapsed, "q2_pass": design_q2["all_constraints_pass"], "q3_pass": design_q3["all_constraints_pass"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
