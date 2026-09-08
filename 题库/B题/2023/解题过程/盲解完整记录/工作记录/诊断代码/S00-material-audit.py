"""Read-only XLSX diagnostics for the material-audit stage.

The script never saves an input workbook and contains no problem-solving,
prediction, fitting, optimization, simulation, or search logic.
"""

from __future__ import annotations

import hashlib
import json
import math
import statistics
import sys
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        if math.isinf(value):
            return "+Infinity" if value > 0 else "-Infinity"
    return value


def bbox(cells: list[tuple[int, int]]) -> str | None:
    if not cells:
        return None
    rows = [row for row, _ in cells]
    cols = [col for _, col in cells]
    return (
        f"{get_column_letter(min(cols))}{min(rows)}:"
        f"{get_column_letter(max(cols))}{max(rows)}"
    )


def analyze_sheet(ws, cached_ws) -> dict[str, Any]:
    content_cells: list[tuple[int, int]] = []
    styled_cells: list[tuple[int, int]] = []
    type_counts: Counter[str] = Counter()
    formulas: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    hyperlinks: list[dict[str, Any]] = []
    non_finite: list[dict[str, Any]] = []
    date_cells: list[dict[str, Any]] = []
    strings: list[dict[str, Any]] = []
    numeric_by_col: dict[int, list[float]] = {}
    style_counts: Counter[str] = Counter()
    number_format_counts: Counter[str] = Counter()

    for row in ws.iter_rows(
        min_row=1,
        max_row=max(ws.max_row, 1),
        min_col=1,
        max_col=max(ws.max_column, 1),
    ):
        for cell in row:
            if cell.has_style:
                styled_cells.append((cell.row, cell.column))
            value = cell.value
            if value is None:
                continue
            content_cells.append((cell.row, cell.column))
            type_counts[cell.data_type] += 1
            style_counts[str(cell.style_id)] += 1
            number_format_counts[cell.number_format] += 1
            if cell.data_type == "f":
                formulas.append(
                    {
                        "cell": cell.coordinate,
                        "formula": value,
                        "cached_value": json_value(cached_ws[cell.coordinate].value),
                    }
                )
            elif cell.data_type == "e":
                errors.append({"cell": cell.coordinate, "value": str(value)})
            if cell.hyperlink:
                hyperlinks.append(
                    {
                        "cell": cell.coordinate,
                        "target": cell.hyperlink.target,
                        "location": cell.hyperlink.location,
                    }
                )
            if isinstance(value, float) and not math.isfinite(value):
                non_finite.append({"cell": cell.coordinate, "value": json_value(value)})
            if cell.is_date or isinstance(value, (datetime, date)):
                date_cells.append({"cell": cell.coordinate, "value": json_value(value)})
            if isinstance(value, str) and not value.startswith("="):
                strings.append({"cell": cell.coordinate, "value": value})
            if (
                isinstance(value, (int, float))
                and not isinstance(value, bool)
                and math.isfinite(float(value))
            ):
                numeric_by_col.setdefault(cell.column, []).append(float(value))

    content_box = bbox(content_cells)
    styled_box = bbox(styled_cells)
    rectangular_cells = 0
    blank_inside_content_box = 0
    sample_rows: list[dict[str, Any]] = []
    duplicate_counter: Counter[tuple[Any, ...]] = Counter()

    if content_cells:
        min_row = min(row for row, _ in content_cells)
        max_row = max(row for row, _ in content_cells)
        min_col = min(col for _, col in content_cells)
        max_col = max(col for _, col in content_cells)
        rectangular_cells = (max_row - min_row + 1) * (max_col - min_col + 1)
        blank_inside_content_box = rectangular_cells - len(content_cells)
        for row_idx in range(min_row, max_row + 1):
            values = [
                ws.cell(row_idx, col_idx).value
                for col_idx in range(min_col, max_col + 1)
            ]
            if any(value is not None for value in values):
                signature = tuple(json.dumps(json_value(v), ensure_ascii=False) for v in values)
                duplicate_counter[signature] += 1
                if len(sample_rows) < 12:
                    sample_rows.append(
                        {
                            "row": row_idx,
                            "values": [json_value(value) for value in values],
                        }
                    )

    duplicate_row_groups = [
        {
            "count": count,
            "values": [json.loads(item) for item in signature],
        }
        for signature, count in duplicate_counter.items()
        if count > 1
    ]

    numeric_columns = {}
    for column, values in numeric_by_col.items():
        numeric_columns[get_column_letter(column)] = {
            "count": len(values),
            "min": min(values),
            "max": max(values),
            "unique_count": len(set(values)),
            "duplicate_value_count": len(values) - len(set(values)),
        }

    merged_ranges = [str(item) for item in ws.merged_cells.ranges]
    hidden_rows = [idx for idx, dim in ws.row_dimensions.items() if dim.hidden]
    hidden_columns = [key for key, dim in ws.column_dimensions.items() if dim.hidden]
    tables = [
        {"name": table.name, "display_name": table.displayName, "ref": table.ref}
        for table in ws.tables.values()
    ]
    unit_like_strings = [
        item
        for item in strings
        if any(token in item["value"] for token in ("米", "海里", "m", "%", "°", "度"))
    ]

    return {
        "title": ws.title,
        "state": ws.sheet_state,
        "declared_dimension": ws.calculate_dimension(),
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "content_bbox": content_box,
        "styled_bbox": styled_box,
        "nonempty_cell_count": len(content_cells),
        "rectangular_cell_count": rectangular_cells,
        "blank_inside_content_bbox": blank_inside_content_box,
        "data_type_counts": dict(sorted(type_counts.items())),
        "nonempty_style_id_counts": dict(sorted(style_counts.items())),
        "nonempty_number_format_counts": dict(sorted(number_format_counts.items())),
        "formula_count": len(formulas),
        "formulas": formulas,
        "error_cells": errors,
        "non_finite_cells": non_finite,
        "date_cells": date_cells,
        "hyperlinks": hyperlinks,
        "merged_ranges": merged_ranges,
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        "auto_filter": ws.auto_filter.ref,
        "tables": tables,
        "print_area": str(ws.print_area) if ws.print_area else None,
        "hidden_rows": hidden_rows,
        "hidden_columns": hidden_columns,
        "sample_nonempty_rows": sample_rows,
        "all_strings": strings if len(strings) <= 80 else strings[:80],
        "all_strings_truncated": len(strings) > 80,
        "unit_like_strings": unit_like_strings,
        "duplicate_full_row_groups": duplicate_row_groups[:20],
        "duplicate_full_row_groups_truncated": len(duplicate_row_groups) > 20,
        "numeric_columns": numeric_columns,
    }


def sequence_diagnostics(values: list[Any]) -> dict[str, Any]:
    numeric = [float(value) for value in values if isinstance(value, (int, float))]
    steps = [numeric[index + 1] - numeric[index] for index in range(len(numeric) - 1)]
    return {
        "cell_count": len(values),
        "numeric_count": len(numeric),
        "blank_count": sum(value is None for value in values),
        "non_numeric_count": sum(
            value is not None and not isinstance(value, (int, float)) for value in values
        ),
        "unique_count": len(set(numeric)),
        "duplicate_count": len(numeric) - len(set(numeric)),
        "first": numeric[0] if numeric else None,
        "last": numeric[-1] if numeric else None,
        "min": min(numeric) if numeric else None,
        "max": max(numeric) if numeric else None,
        "step_min": min(steps) if steps else None,
        "step_max": max(steps) if steps else None,
        "non_increasing_steps": sum(step <= 0 for step in steps),
    }


def material_role_diagnostics(filename: str, ws) -> dict[str, Any]:
    if filename == "附件01.xlsx":
        distances = [ws.cell(1, col).value for col in range(2, 11)]
        output_cells = [
            ws.cell(row, col).value
            for row, col_start in ((2, 2), (3, 2), (4, 3))
            for col in range(col_start, 11)
        ]
        return {
            "observed_role": "problem-1 result template",
            "distance_axis": sequence_diagnostics(distances),
            "prefilled_center_depth_cell": {"cell": "F2", "value": ws["F2"].value},
            "first_overlap_marker": {"cell": "B4", "value": ws["B4"].value},
            "output_cells_checked": len(output_cells),
            "blank_output_cells": sum(value is None for value in output_cells),
        }
    if filename == "附件02.xlsx":
        distances = [ws.cell(2, col).value for col in range(3, 11)]
        angles = [ws.cell(row, 2).value for row in range(3, 11)]
        output_cells = [
            ws.cell(row, col).value for row in range(3, 11) for col in range(3, 11)
        ]
        return {
            "observed_role": "problem-2 result template",
            "distance_axis": sequence_diagnostics(distances),
            "angle_axis": sequence_diagnostics(angles),
            "output_range": "C3:J10",
            "output_cells_checked": len(output_cells),
            "blank_output_cells": sum(value is None for value in output_cells),
        }
    if filename == "附件03.xlsx":
        xs = [ws.cell(2, col).value for col in range(3, 204)]
        ys = [ws.cell(row, 2).value for row in range(3, 254)]
        depths: list[float] = []
        missing_depth_cells: list[str] = []
        non_numeric_depth_cells: list[str] = []
        min_locations: list[str] = []
        max_locations: list[str] = []
        horizontal_jumps: list[tuple[float, str, str]] = []
        vertical_jumps: list[tuple[float, str, str]] = []
        for row in range(3, 254):
            for col in range(3, 204):
                cell = ws.cell(row, col)
                value = cell.value
                if value is None:
                    missing_depth_cells.append(cell.coordinate)
                elif not isinstance(value, (int, float)) or not math.isfinite(float(value)):
                    non_numeric_depth_cells.append(cell.coordinate)
                else:
                    depths.append(float(value))
                    if col < 203:
                        right = ws.cell(row, col + 1)
                        if isinstance(right.value, (int, float)):
                            horizontal_jumps.append(
                                (abs(float(right.value) - float(value)), cell.coordinate, right.coordinate)
                            )
                    if row < 253:
                        below = ws.cell(row + 1, col)
                        if isinstance(below.value, (int, float)):
                            vertical_jumps.append(
                                (abs(float(below.value) - float(value)), cell.coordinate, below.coordinate)
                            )
        if depths:
            minimum = min(depths)
            maximum = max(depths)
            for row in range(3, 254):
                for col in range(3, 204):
                    value = ws.cell(row, col).value
                    if value == minimum:
                        min_locations.append(ws.cell(row, col).coordinate)
                    if value == maximum:
                        max_locations.append(ws.cell(row, col).coordinate)
        else:
            minimum = maximum = None
        max_horizontal = max(horizontal_jumps, default=None)
        max_vertical = max(vertical_jumps, default=None)
        return {
            "observed_role": "bathymetric input grid",
            "x_axis_range": "C2:GU2",
            "x_axis": sequence_diagnostics(xs),
            "y_axis_range": "B3:B253",
            "y_axis": sequence_diagnostics(ys),
            "depth_range": "C3:GU253",
            "expected_depth_cells": 251 * 201,
            "numeric_depth_cells": len(depths),
            "missing_depth_cell_count": len(missing_depth_cells),
            "missing_depth_cells": missing_depth_cells[:20],
            "non_numeric_depth_cell_count": len(non_numeric_depth_cells),
            "non_numeric_depth_cells": non_numeric_depth_cells[:20],
            "finite_depth_count": len(depths),
            "positive_depth_count": sum(value > 0 for value in depths),
            "depth_min": minimum,
            "depth_min_locations": min_locations,
            "depth_max": maximum,
            "depth_max_locations": max_locations,
            "depth_mean": statistics.fmean(depths) if depths else None,
            "depth_median": statistics.median(depths) if depths else None,
            "depth_unique_count": len(set(depths)),
            "depth_duplicate_entry_count": len(depths) - len(set(depths)),
            "values_with_more_than_2_decimals": sum(
                abs(value * 100 - round(value * 100)) > 1e-8 for value in depths
            ),
            "max_horizontal_adjacent_absolute_difference": (
                {"difference": max_horizontal[0], "cells": max_horizontal[1:]}
                if max_horizontal
                else None
            ),
            "max_vertical_adjacent_absolute_difference": (
                {"difference": max_vertical[0], "cells": max_vertical[1:]}
                if max_vertical
                else None
            ),
        }
    return {"observed_role": "unclassified"}


def analyze_workbook(path: Path, root: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(
        path, read_only=False, data_only=False, keep_links=True
    )
    cached = openpyxl.load_workbook(
        path, read_only=False, data_only=True, keep_links=True
    )
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    external_link_parts = [name for name in names if name.startswith("xl/externalLinks/")]
    macro_parts = [name for name in names if name.endswith("vbaProject.bin")]
    return {
        "relative_path": path.relative_to(root).as_posix(),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "format": path.suffix.lower(),
        "sheet_names": workbook.sheetnames,
        "defined_names": [name.name for name in workbook.defined_names.values()],
        "calculation_mode": workbook.calculation.fullCalcOnLoad,
        "external_link_parts": external_link_parts,
        "macro_parts": macro_parts,
        "material_role_diagnostics": material_role_diagnostics(path.name, workbook.active),
        "sheets": [
            analyze_sheet(workbook[name], cached[name]) for name in workbook.sheetnames
        ],
    }


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: S00-material-audit.py WORKSPACE OUTPUT_JSON")
    root = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    attachments = [
        root / "附件" / "附件01.xlsx",
        root / "附件" / "附件02.xlsx",
        root / "附件" / "附件03.xlsx",
    ]
    report = {
        "purpose": "material integrity and workbook-type diagnostics only",
        "input_files_modified": False,
        "workbooks": [analyze_workbook(path, root) for path in attachments],
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
