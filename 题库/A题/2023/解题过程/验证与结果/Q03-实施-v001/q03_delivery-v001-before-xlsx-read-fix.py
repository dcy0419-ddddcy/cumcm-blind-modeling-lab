"""Q3 actual Excel precision freeze and later gated delivery; root owns budget."""
from pathlib import Path
import copy,math,json
import q03_common_v001 as io
HEADER=['吸收塔x坐标 (m)','吸收塔y坐标 (m)','定日镜序号','定日镜宽度 (m)','定日镜高度 (m)','定日镜x坐标 (m)','定日镜y坐标 (m)','定日镜z坐标 (m)']

def read_excel(path,source):
 from openpyxl import load_workbook
 w=load_workbook(path,read_only=True,data_only=True);s=w['完整清单'];rows=list(s.values)
 if list(rows[0])!=HEADER:raise ValueError('XLSX_HEADER_DIFFERS')
 identity=list(w['身份映射'].values());w.close()
 if len(rows)!=len(source['mirrors'])+1 or len(identity)!=len(rows):raise ValueError('XLSX_ROW_COUNT')
 des=copy.deepcopy(source)
 for j,(r,m)in enumerate(zip(rows[1:],des['mirrors'])):
  if r[2]!=j+1 or identity[j+1]!=(j+1,m['mirror_id'],m['position_key'],m['group']):raise ValueError('XLSX_IDENTITY')
  if list(r[:2])!=list(des['tower_xy']):raise ValueError('XLSX_TOWER')
  m.update(width=float(r[3]),height=float(r[4]),x=float(r[5]),y=float(r[6]),z=float(r[7]));m['area']=m['width']*m['height']
 return des

def freeze(name,newname,budget):
 from openpyxl import Workbook
 import q03_engine_v001 as e,q03_design_v001 as d,q03_parallel_v001 as p
 src=io.OUT/'candidates'/name/'design.json';des=io.load(src);des['name']=newname;des['version']='q03-excel-delivery-precision-v001'
 folder=io.OUT/'candidates'/newname;folder.mkdir(parents=True,exist_ok=True)
 xlsx=folder/'候选交付清单-未确认.xlsx'
 if xlsx.exists():raise FileExistsError(xlsx)
 wb=Workbook();s=wb.active;s.title='完整清单';s.append(HEADER);sid=wb.create_sheet('身份映射');sid.append(['导出序号','稳定镜号','位置身份','组号'])
 for j,m in enumerate(des['mirrors']):
  s.append([*des['tower_xy'],j+1,m['width'],m['height'],m['x'],m['y'],m['z']]);sid.append([j+1,m['mirror_id'],m['position_key'],m['group']])
 for sh in wb:
  sh.freeze_panes='A2'
  for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=24
 wb.save(xlsx);wb.close();back=read_excel(xlsx,des)
 changes=[{'index':j,'field':key,'before':a[key],'after':b[key]}for j,(a,b)in enumerate(zip(des['mirrors'],back['mirrors']))for key in ['x','y','width','height','z','area']if a[key]!=b[key]]
 f=e.prepare_design(back,folder,budget)
 original=io.load(io.OUT/'candidates/Q3R000/design.json')
 assert len(back['mirrors'])==len(original['mirrors'])
 assert back['tower_xy']==original['tower_xy']
 assert all((a['x'],a['y'],a['position_key'],a['mirror_id'],a['group'])==(b['x'],b['y'],b['position_key'],b['mirror_id'],b['group'])for a,b in zip(back['mirrors'],original['mirrors']))
 receipt={'candidate':newname,'search_parent':name,'input_source_sha256':io.sha(src),'xlsx':str(xlsx.relative_to(io.OUT)),'xlsx_sha256':io.sha(xlsx),'readback_changes':changes,'readback_geometry':f.validation,'frozen_design_key':f.key,'frozen_design_sha256':io.sha(folder/'design.json'),'fixed_xy_tower_N_identity':True,'execution_sources':p._source_hashes(),'planned_confirmation':{'B':8,'n':256,'levels':[128,256],'namespace':3190,'root_seed':2026090506,'time_count':60,'actual_mirrors':f.n,'candidate_and_baseline_paired':True,'max_n':256,'refinement':'Current supported chain frozen at256 with correlated128 diagnostic prefix. If insufficient, retain failure; no unsupported512 run or seed retry. No tuning with confirmation.'},'work_indicator':'max(2.3646242515927844*whole-batch JK SE, final-half-prefix change, raw alternative shift)+actual-area low survival diameter; pointwise approximate statistical term, no joint/physical guarantee','paired_improvement':'candidate-minus-baseline point and delete-one-batch paired estimates; Udelta=max(t7 paired JK, paired prefix change, paired raw change)+sum of annual candidate and baseline q low-survival diameter bounds. delta_q>Udelta is numerical work evidence, not strict confidence statement.'}
 receipt.update(rules_sha256=io.sha(io.OUT/'搜索规则冻结-v001.json'),baseline_design_sha256=io.sha(io.OUT/'candidates/Q3R000/design.json'),launch_sources={k:io.sha(Path(__file__).parent/k)for k in ['q03_run_v001.py','q03_job_v001.ps1','q03_delivery_v001.py','q03_verify_v001.py']})
 io.save(io.OUT/'最终候选与确认冻结-v001.json',receipt)
 io.save(folder/'export-readback-before-confirmation.json',receipt)
 for nm in [newname,'Q3R000']:
  io.save(io.OUT/'specs'/(nm+'-confirmation.json'),{'name':nm,'tag':'confirmation','B':8,'n':256,'levels':[128,256],'namespace':3190,'mode':'full_confirmation','reserve_after':90,'changes':{},'rules':'搜索规则冻结-v001.json','freeze':'最终候选与确认冻结-v001.json'})
 return receipt

def deliver(budget):
 import csv,shutil,numpy as np
 import q03_verify_v001 as verifier
 ans=verifier.run(budget)
 if not(ans['all_pass'] and ans['formal_ready'] and ans['improvement_supported']):return ans
 freeze=io.load(io.OUT/'最终候选与确认冻结-v001.json');name=freeze['candidate'];folder=io.OUT/'candidates'/name
 des=io.load(folder/'design.json');summary=io.load(folder/'confirmation/summary.json');point=np.asarray(summary['point']);U=np.asarray(summary['conservative_work_indicator'])
 out=io.OUT/'delivery';out.mkdir(exist_ok=True)
 src=io.OUT/freeze['xlsx'];target=out/'result3.xlsx'
 if target.exists()and io.sha(target)!=io.sha(src):raise FileExistsError('Existing official workbook differs')
 if not target.exists():shutil.copyfile(src,target)
 assert read_excel(target,des)==des
 tables={
 'table1':{'header':['日期','平均光学效率','平均余弦效率','平均阴影遮挡效率','平均截断效率','单位面积镜面平均输出热功率 (kW/m²)'],'rows':[[str(m+1)+'月21日',*point[m,:4],point[m,5]]for m in range(12)]},
 'table2':{'header':['年平均光学效率','年平均余弦效率','年平均阴影遮挡效率','年平均截断效率','年平均输出热功率 (MW)','单位面积镜面年平均输出热功率 (kW/m²)'],'rows':[[*point[-1,:4],point[-1,4]/1000,point[-1,5]]]},
 'table3':{'header':['吸收塔位置坐标 (m)','定日镜尺寸（宽×高）(m)','定日镜安装高度 (m)','定日镜总面数','定日镜总面积 (m²)'],'rows':[[str(tuple(des['tower_xy'])),'','',len(des['mirrors']),sum(r['area']for r in des['mirrors'])]]}}
 for k,v in tables.items():
  path=out/(k+'.csv')
  with path.open('w',newline='',encoding='utf-8-sig')as handle:
   wr=csv.writer(handle);wr.writerow(v['header']);wr.writerows(v['rows'])
  with path.open(encoding='utf-8-sig',newline='')as handle:rr=list(csv.reader(handle))
  assert rr==[[str(x)for x in v['header']]]+[[str(x)for x in row]for row in v['rows']]
 result={'version':'q03-results-v001','selected_design':name,'tables':tables,'point':point,'work_indicator':U,'annual_point_internal_units':point[-1],'annual_work_indicator_internal_units':U[-1],'area_weighted':summary['area_weighted_efficiencies'],'energy_total_ratios':summary['energy_total_ratios'],'rating_margin_kw':point[-1,4]-U[-1,4]-60000,'N':len(des['mirrors']),'total_area':sum(r['area']for r in des['mirrors']),'groups':[{'group':g,'n':sum(r['group']==g for r in des['mirrors']),'width':next(r['width']for r in des['mirrors']if r['group']==g),'height':next(r['height']for r in des['mirrors']if r['group']==g),'z':next(r['z']for r in des['mirrors']if r['group']==g)}for g in sorted({r['group']for r in des['mirrors']})]}
 io.save(out/'结果数据-v001.json',result)
 paths={'design_file':str((folder/'design.json').relative_to(io.OUT)),'candidate_confirmation_summary':str((folder/'confirmation/summary.json').relative_to(io.OUT)),'baseline_confirmation_summary':'candidates/Q3R000/confirmation/summary.json','comparison_report':'delivery/paired-comparison-v001.json','excel_file':'delivery/result3.xlsx','rebuild_report':'delivery/独立重建核验-v001.json','search_trace':'scores.jsonl','result_data':'delivery/结果数据-v001.json'}
 manifest={'status':'CONFIRMED_IMPROVEMENT','selected_design':name,**paths,'sha256':{v:io.sha(io.OUT/v)for v in paths.values()},'physical_error':'unquantified','global_optimum_claim':False}
 io.save(out/'report.json',manifest)
 return {**ans,'delivered':True,'name':name,'annual':point[-1],'U':U[-1],'rating_margin_kw':result['rating_margin_kw']}

