"""只读材料诊断。仅处理此单题目录；输出保存在工作记录内。"""
from pathlib import Path
import hashlib
import json
import sys
import platform
import importlib.metadata
import math
import csv
import zipfile
from collections import Counter
from datetime import datetime
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '工作记录' / '诊断结果'
OUT.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding='utf-8')

def dump(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

def main():
    inputs = [ROOT / n for n in ['AGENTS.md','README.md','总控提示词.md','第一步-材料审计与题目理解.md','题目.pdf']]
    inputs += sorted((ROOT / '附件').glob('*'))
    inventory = []
    for p in inputs:
        if p.is_file():
            if not p.resolve().is_relative_to(ROOT):
                raise RuntimeError('材料路径越界')
            inventory.append({'path':p.relative_to(ROOT).as_posix(),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    dump('原始材料清单-v001.json', inventory)
    dump('环境-v001.json', {'python':sys.version, 'platform':platform.platform(), 'packages':{m:importlib.metadata.version(m) for m in ['pypdf','pdfplumber','openpyxl','Pillow']},'runtime':'Codex bundled local Python','random_seed':'不适用；无随机过程','timestamp':datetime.now().astimezone().isoformat(),'method_library_exists':(ROOT/'方法库').exists()})
    reader = PdfReader(ROOT / '题目.pdf')
    pages=[]
    text=[]
    for i, page in enumerate(reader.pages, 1):
        s=page.extract_text() or ''
        pages.append({'page':i, 'size_points':[float(x) for x in page.mediabox], 'characters':len(s)})
        text.append(f'--- PDF 第 {i} 页 ---\n{s}')
    dump('题目结构-v001.json',{'page_count':len(pages),'encrypted':reader.is_encrypted,'pages':pages})
    (OUT/'题目逐页提取-v001.txt').write_text('\n\n'.join(text),encoding='utf-8')
    books=[]
    for p in sorted((ROOT/'附件').glob('*.xlsx')):
        wb=load_workbook(p,read_only=False,data_only=False,keep_links=False)
        book={'file':p.relative_to(ROOT).as_posix(),'sheets':[]}
        for ws in wb:
            book['sheets'].append({'name':ws.title,'state':ws.sheet_state,'dimensions':ws.calculate_dimension(),'max_row':ws.max_row,'max_column':ws.max_column,'merged_ranges':[str(x) for x in ws.merged_cells.ranges],'head_rows':list(ws.iter_rows(min_row=1,max_row=min(8,ws.max_row),values_only=True))})
        wb.close()
        books.append(book)
    dump('工作簿结构-v001.json', books)
    print(json.dumps({'inventory':inventory,'pdf':pages,'workbooks':books},ensure_ascii=False,indent=2,default=str))
    audit_workbooks()

def audit_workbooks():
    reports=[]
    for p in sorted((ROOT/'附件').glob('*.xlsx')):
        with zipfile.ZipFile(p) as z:
            names=z.namelist()
            zip_check=z.testzip()
            package={'crc_error_member':zip_check,'vba_present':any('vbaProject' in n for n in names),'external_link_parts':[n for n in names if n.startswith('xl/externalLinks/')],'embedding_parts':[n for n in names if n.startswith('xl/embeddings/')]}
        wb=load_workbook(p,read_only=False,data_only=False,keep_links=False)
        report={'file':p.relative_to(ROOT).as_posix(),'package':package,'sheets':[]}
        for ws in wb:
            rows=list(ws.iter_rows(values_only=True))
            cells=list(ws.iter_rows())
            sr={'name':ws.title,'range':ws.calculate_dimension(),'headers':rows[0],'data_rows':len(rows)-1,'empty_cells_below_header':sum(v is None for row in rows[1:] for v in row),'fully_empty_data_rows':[i for i,row in enumerate(rows[1:],2) if all(v is None for v in row)],'formula_cells':[c.coordinate for row in cells for c in row if c.data_type=='f'],'error_cells':[c.coordinate for row in cells for c in row if c.data_type=='e'],'comment_cells':[c.coordinate for row in cells for c in row if c.comment],'hyperlink_cells':[c.coordinate for row in cells for c in row if c.hyperlink],'hidden_rows':[i for i,d in ws.row_dimensions.items() if d.hidden],'hidden_columns':[i for i,d in ws.column_dimensions.items() if d.hidden],'auto_filter':ws.auto_filter.ref,'table_names':list(ws.tables),'image_count':len(ws._images),'chart_count':len(ws._charts),'column_type_counts':{str(k+1):dict(Counter(type(row[k]).__name__ for row in rows[1:])) for k in range(ws.max_column)}}
            if p.name=='附件03.xlsx':
                records=[]
                invalid=[]
                for n,row in enumerate(rows[1:],2):
                    if len(row)==2 and all(isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v) for v in row):
                        records.append((n,float(row[0]),float(row[1])))
                    else: invalid.append({'excel_row':n,'values':row})
                counts=Counter((x,y) for _,x,y in records)
                sr['invalid_coordinate_rows']=invalid
                sr['duplicate_coordinate_groups']=[{'point':xy,'count':n} for xy,n in counts.items() if n>1]
                sr['x_range_m']=[min(x for _,x,y in records),max(x for _,x,y in records)]
                sr['y_range_m']=[min(y for _,x,y in records),max(y for _,x,y in records)]
                radii=[math.hypot(x,y) for _,x,y in records]
                sr['radius_range_m']=[min(radii),max(radii)]
                sr['radius_below_100_rows']=[n for (n,x,y),r in zip(records,radii) if r<100]
                sr['radius_above_350_rows']=[n for (n,x,y),r in zip(records,radii) if r>350]
                sr['at_radial_boundary_rows']=[n for (n,x,y),r in zip(records,radii) if r in (100,350)]
                nearest=[float('inf')]*len(records)
                near_row=[None]*len(records)
                flagged=[]
                for i,(ni,xi,yi) in enumerate(records):
                    for j in range(i+1,len(records)):
                        nj,xj,yj=records[j]
                        d=math.hypot(xi-xj,yi-yj)
                        if d<nearest[i]:nearest[i],near_row[i]=d,nj
                        if d<nearest[j]:nearest[j],near_row[j]=d,ni
                        if d<=11:flagged.append({'excel_row_1':ni,'excel_row_2':nj,'distance_m':d,'shortfall_from_11_m':11-d})
                sr['coordinate_count']=len(records)
                sr['distance_check']='所有无序点对的平面欧氏距离；11 m 仅是第1问6 m镜宽加题面5 m的审计参照，未自行设置容差。'
                sr['min_pair_distance_m']=min(nearest)
                sr['max_nearest_neighbor_distance_m']=max(nearest)
                sr['pairs_at_or_below_11_m']=len(flagged)
                sr['pairs_strictly_below_11_m']=sum(v['distance_m']<11 for v in flagged)
                sr['pairs_at_or_below_11_details']=flagged
                sr['coordinate_rounding_check']='所有数值是否均落在0.001 m网格（不据此断言测量精度）'
                sr['all_coordinates_on_0_001_m_grid']=all(abs(v*1000-round(v*1000))<1e-8 for _,x,y in records for v in (x,y))
                with (OUT/'坐标诊断明细-v001.csv').open('w',encoding='utf-8-sig',newline='') as f:
                    writer=csv.writer(f)
                    writer.writerow(['excel_row','x_m','y_m','radius_m','nearest_excel_row','nearest_distance_m'])
                    for rec,r,nr,nd in zip(records,radii,near_row,nearest):writer.writerow([*rec,r,nr,nd])
                with (OUT/'间距边界记录-v001.csv').open('w',encoding='utf-8-sig',newline='') as f:
                    writer=csv.DictWriter(f,fieldnames=['excel_row_1','excel_row_2','distance_m','shortfall_from_11_m'])
                    writer.writeheader();writer.writerows(flagged)
            report['sheets'].append(sr)
        wb.close()
        reports.append(report)
    dump('附件数据诊断-v001.json',reports)
    compact=json.loads(json.dumps(reports,ensure_ascii=False,default=str))
    for book in compact:
        for sheet in book['sheets']:
            if 'pairs_at_or_below_11_details' in sheet:
                sheet['pairs_at_or_below_11_details']=sorted(sheet['pairs_at_or_below_11_details'],key=lambda x:x['distance_m'])[:10]
    print(json.dumps(compact,ensure_ascii=False,indent=2))

if __name__ == '__main__':
    main()
