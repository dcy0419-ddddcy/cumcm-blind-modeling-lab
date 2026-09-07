from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

W = Path(r"C:\Users\admin\OneDrive\Desktop\26国赛-单题盲解区\A-P8C3")
P = W / "工作记录/揭晓后对照/首次正确性核查-v001"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_values(path: Path):
    wb = load_workbook(path, read_only=True, data_only=False)
    return [(ws.title, [[c.value for c in row] for row in ws.iter_rows()]) for ws in wb.worksheets]


official_pdf = P / "官方来源/A题.pdf"
local_pdf = W / "题目.pdf"
official_pages = [p.extract_text() or "" for p in PdfReader(official_pdf).pages]
local_pages = [p.extract_text() or "" for p in PdfReader(local_pdf).pages]
excel_pairs = [
    (P / "官方来源/附件.xlsx", W / "附件/附件03.xlsx", "coordinates"),
    (P / "官方来源/result2.xlsx", W / "附件/附件01.xlsx", "result2_template"),
    (P / "官方来源/result3.xlsx", W / "附件/附件02.xlsx", "result3_template"),
]
out = {
    "official_pdf_sha256": sha(official_pdf),
    "local_pdf_sha256": sha(local_pdf),
    "official_pdf_pages": len(official_pages),
    "local_pdf_pages": len(local_pages),
    "pdf_page_text_equal": [a == b for a, b in zip(official_pages, local_pages)],
    "pdf_note": "Page 1 differs only because the local anonymous copy omits the competition/year heading; pages 2-4 extracted text equal. Visual page renders were also checked.",
    "workbooks": [],
}
for official, local, identity in excel_pairs:
    ov = workbook_values(official)
    lv = workbook_values(local)
    out["workbooks"].append({
        "identity": identity,
        "official_sha256": sha(official),
        "local_sha256": sha(local),
        "all_cell_values_equal": ov == lv,
        "official_sheets": [x[0] for x in ov],
        "local_sheets": [x[0] for x in lv],
    })
(P / "官方身份核对-v001.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False))
