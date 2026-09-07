from __future__ import annotations
import hashlib, json, re, subprocess
from pathlib import Path
from datetime import datetime
REPO=Path(__file__).resolve().parents[5]
BASE=REPO/"题库/A题/2023"
OUT=BASE/"解题过程/管理员归档审计/提交前核验-v001.json"
SRC=Path(r"C:\Users\admin\OneDrive\Desktop\26国赛-单题盲解区\A-P8C3")
EXPECTED_MANIFEST="ed23098187e4479ab250744f6cc35fc84e1ec5bb25afe7e6b4ac21851c0da357"
EXPECTED_SUMS="9e6b950d124e16454aaf1715d21d2be183432e2703f4b92ff9a10d0d01820439"
def sha(p):
 h=hashlib.sha256()
 with p.open("rb") as f:
  for b in iter(lambda:f.read(1024*1024),b""):h.update(b)
 return h.hexdigest()
def git(*args):
 return subprocess.run(["git","-c","core.quotepath=false",*args],cwd=REPO,text=True,encoding="utf-8",stdout=subprocess.PIPE,stderr=subprocess.PIPE,check=True).stdout.splitlines()
def links(p):
 text=p.read_text(encoding="utf-8-sig"); bad=[]
 for t in re.findall(r"\[[^\]]+\]\(([^)]+)\)",text):
  t=t.strip().strip("<>").split("#",1)[0]
  if not t or re.match(r"^[a-zA-Z]+://",t):continue
  q=(p.parent/t).resolve()
  if not q.exists():bad.append(t)
 return bad
def secret_scan(paths):
 # Deliberately assembled so the scanner does not match its own source.
 pats={
  "github_token":re.compile("gh"+"[pousr]_[A-Za-z0-9]{20,}"),
  "private_key":re.compile("BEGIN "+"(?:RSA |EC |OPENSSH )?PRIVATE KEY"),
  "generic_secret":re.compile(r"(?i)(api[_-]?key|access[_-]?token|client[_-]?secret)\s*[:=]\s*[\"']?[A-Za-z0-9_\-]{16,}")
 }
 hits=[]
 for rel in paths:
  p=REPO/rel
  if not p.is_file() or p.stat().st_size>5_000_000 or p.suffix.lower() in {".pdf",".xlsx",".png",".jpg",".jpeg",".npz"}:continue
  try:text=p.read_text(encoding="utf-8")
  except UnicodeDecodeError:continue
  for kind,rx in pats.items():
   if rx.search(text):hits.append({"path":rel,"type":kind})
 return hits
def main():
 staged=git("diff","--cached","--name-only")
 readmes=[
  BASE/"README.md",BASE/"解题过程/00-任务状态.md",BASE/"解题过程/01-题目理解.md",
  BASE/"解题过程/02-逐问建模.md",BASE/"解题过程/03-求解与验证.md",
  BASE/"解题过程/04-结果与论文.md",BASE/"解题过程/05-优秀论文复盘.md",
  BASE/"解题过程/06-更新记录.md",BASE/"解题过程/代码与配置/README.md",
  BASE/"解题过程/验证与结果/README.md",BASE/"解题过程/失败与修订/README.md",
  BASE/"解题过程/冻结版本/README.md",BASE/"解题过程/归档映射.md",
  BASE/"揭晓后参考/README.md",BASE/"揭晓后参考/官方来源链接.md",
  BASE/"揭晓后参考/方法比较与创新分析/README.md",
  REPO/"模型方法库/实践候选/2023-A/README.md",REPO/"总经验/候选/2023-A/README.md"
 ]
 broken={str(p.relative_to(REPO)):links(p) for p in readmes if p.exists() and links(p)}
 pdfs=[
  BASE/"解题过程/论文/全题整合/完整论文-v001.pdf",
  BASE/"解题过程/论文/LaTeX/第1问-模型建立与求解-v001.pdf",
  BASE/"解题过程/论文/LaTeX/第2问-模型建立与求解-v002.pdf",
  BASE/"解题过程/论文/LaTeX/第3问-模型建立与求解-v001.pdf"
 ]
 pdf_info={}
 try:
  from pypdf import PdfReader
  for p in pdfs:
   rel=str(p.relative_to(REPO))
   try:pdf_info[rel]={"exists":p.is_file(),"pages":len(PdfReader(str(p)).pages) if p.is_file() else 0}
   except Exception as e:pdf_info[rel]={"exists":p.is_file(),"pages":0,"error":type(e).__name__+":"+str(e)}
 except Exception as e:
  for p in pdfs:pdf_info[str(p.relative_to(REPO))]={"exists":p.is_file(),"pages":0,"error":"pypdf:"+type(e).__name__+":"+str(e)}
 xlsx_info={}
 try:
  from openpyxl import load_workbook
  for name in ["result2.xlsx","result3.xlsx"]:
   p=BASE/"解题过程/验证与结果/正式结果"/name
   wb=load_workbook(p,read_only=True,data_only=False)
   xlsx_info[name]={"sheets":wb.sheetnames,"sha256":sha(p),"source_sha256":sha(SRC/("工作记录/诊断结果/Q02-修复-v001/result2.xlsx" if name=="result2.xlsx" else "工作记录/诊断结果/Q03-实施-v001/delivery/result3.xlsx"))}
   wb.close()
 except Exception as e:xlsx_info={"error":type(e).__name__+":"+str(e)}
 freeze_pkg=SRC/"工作记录/冻结候选/盲解-v001"
 freeze={"manifest_sha256":sha(freeze_pkg/"文件清单-v001.json"),"sums_sha256":sha(freeze_pkg/"SHA256SUMS.txt")}
 freeze["passed"]=freeze["manifest_sha256"]==EXPECTED_MANIFEST and freeze["sums_sha256"]==EXPECTED_SUMS
 experience=(BASE/"经验.md").read_text(encoding="utf-8-sig")
 missing_e=[f"E{i:03d}" for i in range(1,31) if f"E{i:03d}" not in experience]
 thirdparty=[p for p in staged if re.search(r"(参考论文[/\\].*(page-|\.jpg$|\.html$)|官方来源[/\\].*\.(rar|xlsx|pdf|png)$)",p,re.I)]
 allowed=("题库/A题/2023/","模型方法库/实践候选/2023-A/","总经验/")
 out={
  "checked_at":datetime.now().astimezone().isoformat(),
  "branch":git("branch","--show-current")[0],
  "head":git("rev-parse","HEAD")[0],
  "origin":git("remote","get-url","origin")[0],
  "staged_files":len(staged),
  "staged_bytes":sum((REPO/p).stat().st_size for p in staged if (REPO/p).is_file()),
  "largest_20":sorted([{"path":p,"bytes":(REPO/p).stat().st_size} for p in staged if (REPO/p).is_file()],key=lambda x:x["bytes"],reverse=True)[:20],
  "over_100MiB":[p for p in staged if (REPO/p).is_file() and (REPO/p).stat().st_size>100*1024*1024],
  "outside_allowed_prefix":[p for p in staged if not p.replace("\\","/").startswith(allowed)],
  "third_party_fulltext_or_bulk_pages_staged":thirdparty,
  "secret_scan_hits":secret_scan(staged),
  "broken_new_readme_links":broken,
  "pdf_info":pdf_info,
  "xlsx_info":xlsx_info,
  "xlsx_sources_match":all(v.get("sha256")==v.get("source_sha256") for v in xlsx_info.values()) if "error" not in xlsx_info else False,
  "missing_E001_E030":missing_e,
  "freeze":freeze,
 }
 out["passed"]=(out["branch"]=="codex/archive-2023-a" and not out["over_100MiB"] and not out["outside_allowed_prefix"] and not thirdparty and not out["secret_scan_hits"] and not broken and all(v.get("pages",0)>0 for v in pdf_info.values()) and out["xlsx_sources_match"] and not missing_e and freeze["passed"])
 OUT.parent.mkdir(parents=True,exist_ok=True)
 OUT.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding="utf-8")
 print(json.dumps({k:out[k] for k in ["staged_files","staged_bytes","over_100MiB","outside_allowed_prefix","third_party_fulltext_or_bulk_pages_staged","secret_scan_hits","broken_new_readme_links","pdf_info","xlsx_sources_match","missing_E001_E030","freeze","passed"]},ensure_ascii=True))
if __name__=="__main__":main()

