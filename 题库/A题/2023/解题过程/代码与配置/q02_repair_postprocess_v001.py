"""Repair-round independent U reconstruction and gated official exports.
No new rays, optimization, control changes or manuscript edits.  Numerical
imports occur after the standalone cumulative confirmation budget starts.
The official workbook is a byte-identical copy of the frozen confirmed input.
"""
from __future__ import annotations
import time
START = time.perf_counter()
from pathlib import Path
import sys
import json
import csv
import io as string_io
import math
import traceback
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import q02_repair_common_v001 as io

LABELS = ['optical', 'cosine', 'shadow_blocking', 'truncation', 'power_kw', 'unit_power_kw_m2']
OFFICIAL_HEADER = ['吸收塔x坐标 (m)', '吸收塔y坐标 (m)', '定日镜序号', '定日镜宽度 (m)',
                   '定日镜高度 (m)', '定日镜x坐标 (m)', '定日镜y坐标 (m)', '定日镜z坐标 (m)']
T7 = 2.3646242515927844


def write_once(path, data):
    """Idempotent output commit; preserve different existing output bytes."""
    path = Path(path)
    if path.exists():
        if path.read_bytes() != data:
            raise FileExistsError('Different existing delivery bytes: ' + str(path))
        return
    temporary = path.with_suffix(path.suffix + '.part')
    temporary.write_bytes(data)
    io.replace(temporary, path)


def run(budget, report):
    import numpy as np
    import q02_compact_v001 as compact
    O = io.OUT
    S = O / 'frozen' / 'confirmation'
    freeze_path = O / '最终拟提交候选冻结.json'
    freeze = io.load(freeze_path)
    config_path = (O / freeze['config_file']).resolve()
    if config_path.parent != O.resolve():
        raise ValueError('Configuration must remain in this repair output directory')
    cfg = io.load(config_path)
    su = io.load(S / 'summary.json')
    binding = io.load(S / 'binding.json')
    plan = freeze['confirmation_config']
    verification_path = O / '独立重建核验.json'
    verification = io.load(verification_path) if verification_path.exists() else {}
    conclusion_path = O / '独立确认结论.json'
    conclusion = io.load(conclusion_path) if conclusion_path.exists() else {}
    design, receipt = compact.read_compact(O / 'frozen' / 'design.bundle.json', freeze['bundle_sha256'])
    N, B, n = design.n, 8, 256
    levels = list(plan['levels'])
    if (plan['B'], plan['n'], levels) != (B, n, [128, 256]):
        raise ValueError('Expected unchanged fixed B8 n256 half-prefix confirmation')
    if plan != cfg['confirmation'] or binding['namespace'] != plan['namespace'] or plan['namespace'] in cfg['search_namespaces']:
        raise ValueError('Confirmation plan or namespace binding mismatch')
    if binding['root_seed'] != cfg['root_seed'] or binding['design_key'] != freeze['design_key']:
        raise ValueError('Root seed or frozen design binding mismatch')
    if binding['n_mirrors'] != N or binding['name'] != design.name or design.name != freeze['name']:
        raise ValueError('Actual population or candidate identity mismatch')
    if (binding['B'], binding['n'], binding['levels']) != (B, n, levels):
        raise ValueError('Statistics plan differs from frozen confirmation')
    expected_times = [[m, h] for m in range(1, 13) for h in (9., 10.5, 12., 13.5, 15.)]
    if binding['times'] != expected_times:
        raise ValueError('Confirmation must contain the fixed 60 times in standard order')
    threshold = int(plan['low_survival_threshold'])
    summary_sha = io.sha(S / 'summary.json')
    freeze_sha = io.sha(freeze_path)
    binding_sha = io.sha(S / 'binding.json')
    checks = report['checks']
    def check(name, condition, evidence=None):
        checks[name] = {'passed': bool(condition), 'evidence': io.convert(evidence)}
        return bool(condition)
    def compare(name, actual, expected):
        actual, expected = np.asarray(actual, float), np.asarray(expected, float)
        if actual.shape != expected.shape:
            return check(name, False, {'actual_shape': actual.shape, 'expected_shape': expected.shape})
        af, ef = np.isfinite(actual), np.isfinite(expected)
        both = af & ef
        error = float(np.max(np.abs(actual[both] - expected[both]))) if np.any(both) else None
        equal = np.array_equal(af, ef) and np.allclose(actual, expected, atol=2e-9, rtol=2e-9, equal_nan=True)
        return check(name, equal, {'max_abs_difference': error, 'na_pattern_equal': np.array_equal(af, ef),
                                  'actual_na_count': int((~af).sum()), 'atol': 2e-9, 'rtol': 2e-9})
    def one(s, counts, unknown_weight, co, ta, dni, areas, raw=False, sample_count=2048):
        # A sampled zero has no true-zero certificate in this execution path.
        finite = np.all(np.isfinite(s), axis=1)
        known = (counts[:, 4] == 0) & (unknown_weight == 0)
        primary = finite & known & (s[:, 0] > 0)
        sb_good = primary & (s[:, 1] > 0)
        total_good = sb_good & (s[:, 2] > 0)
        eta, sb, trunc = (np.full(N, np.nan) for _ in range(3))
        if raw:
            eta[total_good] = .92 * ta[total_good] * s[total_good, 2] / sample_count
            sb[sb_good] = s[sb_good, 1] / sample_count / co[sb_good]
        else:
            eta[total_good] = .92 * co[total_good] * ta[total_good] * s[total_good, 2] / s[total_good, 0]
            sb[sb_good] = s[sb_good, 1] / s[sb_good, 0]
        trunc[total_good] = s[total_good, 2] / s[total_good, 1]
        power = dni * np.sum(areas * eta)
        return np.array([eta.mean(), co.mean(), sb.mean(), trunc.mean(), power, power / areas.sum()])
    def rows(values):
        values = np.asarray(values, float)
        return np.vstack([values.reshape(12, 5, 6).mean(axis=1), values.mean(axis=0)[None, :]])

    certificate_codes = np.asarray(su.get('certificate_codes'))
    check('no_unhandled_certificates', certificate_codes.shape == (60, N) and bool(np.all(certificate_codes == 0)),
          'This reconstruction supports the current no-certificate confirmation path; it never infers true zero.')
    time_rows, prefix_rows, raw_rows, leave_rows, low_rows = [], [], [], [], []
    chunk_hashes = []
    for ti in range(60):
        budget.guard(15)
        p = S / ('time-%02d.npz' % ti)
        meta = io.load(p.with_suffix('.json'))
        digest = io.sha(p)
        check('chunk_hash_%02d' % ti, digest == meta['sha256'])
        check('chunk_identity_%02d' % ti,
              [meta['month'], meta['hour']] == expected_times[ti] and meta['design_key'] == freeze['design_key']
              and meta['combinations'] == N and meta['unique_source_samples'] == B * n * N)
        chunk_hashes.append(digest)
        with np.load(p, allow_pickle=False) as x:
            ss, cc, uw = x['sums'], x['counts'], x['unknown_weights']
            co, ta, dni, areas = x['cosine'], x['tau'], float(x['dni']), x['areas']
            if ss.shape != (2, B, N, 3) or cc.shape != (2, B, N, 6) or uw.shape != (2, B, N):
                raise ValueError('Raw sufficient-statistics shape mismatch at time ' + str(ti))
            if co.shape != (N,) or ta.shape != (N,) or areas.shape != (N,):
                raise ValueError('Object parameter shape mismatch at time ' + str(ti))
            check('areas_%02d' % ti, bool(np.array_equal(areas, np.full(N, design.width * design.height))))
            s, c, u = ss[-1].sum(axis=0), cc[-1].sum(axis=0), uw[-1].sum(axis=0)
            time_rows.append(one(s, c, u, co, ta, dni, areas))
            prefix_rows.append(one(ss[0].sum(axis=0), cc[0].sum(axis=0), uw[0].sum(axis=0), co, ta, dni, areas))
            current_leave = []
            for batch in range(B):
                keep = [j for j in range(B) if j != batch]
                # Re-sum retained batches directly; avoid cancellation in pool-minus-batch.
                current_leave.append(one(ss[-1, keep].sum(axis=0), cc[-1, keep].sum(axis=0),
                                         uw[-1, keep].sum(axis=0), co, ta, dni, areas))
            leave_rows.append(current_leave)
            raw_rows.append(one(s, c, u, co, ta, dni, areas, raw=True, sample_count=B * n))
            low = c[:, 2] < threshold
            factor = .92 * co * ta
            diameter_power = dni * np.sum(areas[low] * factor[low])
            low_rows.append([factor[low].sum() / N, 0., low.sum() / N, low.sum() / N,
                             diameter_power, diameter_power / areas.sum()])
        if ti % 5 == 4:
            budget.tick()
    budget.guard(20)
    point, prefix, raw = rows(time_rows), rows(prefix_rows), rows(raw_rows)
    leave_time = np.asarray(leave_rows)
    leave = np.asarray([rows(leave_time[:, batch, :]) for batch in range(B)])
    pseudo = B * point - (B - 1) * leave
    se = np.std(pseudo, axis=0, ddof=1) / math.sqrt(B)
    halfwidth = T7 * se
    prefix_change, raw_change = np.abs(point - prefix), np.abs(point - raw)
    base_u = np.maximum.reduce([halfwidth, prefix_change, raw_change])
    low_bound = rows(low_rows)
    U = base_u + low_bound
    pairs = [('point', point, 'point'), ('leave', leave, 'leave_one_batch_estimates'),
             ('pseudo_values', pseudo, 'pseudo_values'), ('prefix', prefix, 'prefix_half_point'),
             ('raw', raw, 'raw_integral_alternative'), ('jackknife_se', se, 'jackknife_se'),
             ('t7_halfwidth', halfwidth, 'approx_pointwise_t7_95_halfwidth'),
             ('base_U', base_u, 'base_numeric_work_indicator'), ('U', U, 'conservative_work_indicator')]
    for label, actual, key in pairs:
        compare(label, actual, su[key])
    compare('prefix_change', prefix_change, su['absolute_changes']['prefix_half'])
    compare('raw_change', raw_change, su['absolute_changes']['raw_integral'])
    compare('low_count_diameter', low_bound, su['low_survival']['effect_diameter_bound'])
    check('low_count_threshold', su['low_survival']['threshold_exclusive'] == threshold)
    check('required_work_values_defined', bool(np.all(np.isfinite(point)) and np.all(np.isfinite(U))
          and np.all(np.isfinite(prefix)) and np.all(np.isfinite(raw)) and np.all(np.isfinite(leave))))
    efficiency_target, power_target = float(plan['efficiency_abs']), float(plan['power_relative'])
    with np.errstate(divide='ignore', invalid='ignore'):
        relative = U[:, 4:] / np.abs(point[:, 4:])
    targets = np.concatenate([np.isfinite(U[:, :4]) & (U[:, :4] <= efficiency_target),
                              np.isfinite(relative) & (np.abs(point[:, 4:]) > 0) & (relative <= power_target)], axis=1)
    lower = point[-1, 4] - U[-1, 4]
    compare('precision_flags', targets, su['target_pass'])
    compare('rating_lower_kw', [lower], [su['rated_power_lower_work_value_kw']])
    check('precision_aggregate', bool(targets.all()) == su['all_table_items_precision_met'])
    check('rated_power_flag', bool(math.isfinite(float(lower)) and lower >= 60000.) == su['rated_power_met'])
    report.update({'name': freeze['name'], 'design_key': freeze['design_key'],
        'confirmation_summary_sha256': summary_sha, 'frozen_candidate_sha256': freeze_sha,
        'confirmation_binding_sha256': binding_sha, 'actual_n': N, 'actual_total_area_m2': design.total_area,
        'point': point, 'prefix_half_point': prefix, 'raw_integral_alternative': raw,
        'leave_one_batch_estimates': leave, 'pseudo_values': pseudo, 'jackknife_se': se,
        'approx_pointwise_t7_95_halfwidth': halfwidth, 'base_numeric_work_indicator': base_u,
        'low_survival_effect_diameter_bound': low_bound, 'conservative_work_indicator': U,
        'rated_lower_kw': lower, 'target_pass': targets, 'new_optical_rays': 0,
        'scope': 'Independent arithmetic of saved statistics; approximate pointwise numerical work indicator, not physical validation or simultaneous confidence.'})
    report['all_pass'] = all(v['passed'] for v in checks.values())
    expected_verified_hashes = {'summary.json': summary_sha, '最终拟提交候选冻结.json': freeze_sha,
                               'binding.json': binding_sha, config_path.name: io.sha(config_path)}
    verification_binding_ok = all(verification.get('source_hashes', {}).get(k) == v for k, v in expected_verified_hashes.items())
    verified_chunks = verification.get('chunks', [])
    verification_binding_ok = verification_binding_ok and len(verified_chunks) == 60 and all(
        row.get('time_index') == ti and row.get('npz_sha256') == chunk_hashes[ti] for ti, row in enumerate(verified_chunks))
    gates = {'confirmation_summary_PASS': su.get('formal_decision') == 'PASS',
             'confirmation_conclusion_PASS': conclusion.get('decision') == 'PASS' and conclusion.get('name') == freeze['name']
                 and conclusion.get('design_key') == freeze['design_key'],
             'independent_point_rebuild_PASS': verification.get('all_pass') is True,
             'independent_point_rebuild_inputs_unchanged': verification_binding_ok,
             'independent_work_rebuild_PASS': report['all_pass'],
             'all_required_metrics_defined': bool(np.all(np.isfinite(point)) and np.all(np.isfinite(U))),
             'reconstructed_precision_and_rating_PASS': bool(targets.all() and math.isfinite(float(lower)) and lower >= 60000.)}
    report['export_gates'] = gates
    report['export_allowed'] = all(gates.values())
    report['exports_created'] = False
    if not report['export_allowed']:
        report['export_withheld_reasons'] = [key for key, passed in gates.items() if not passed]
        return
    budget.guard(30)
    frozen_excel = (O / freeze['excel_file']).resolve()
    if frozen_excel.parent != (O / 'frozen').resolve():
        raise ValueError('Frozen workbook path outside frozen output directory')
    if io.sha(frozen_excel) != freeze['excel_sha256']:
        raise ValueError('Frozen workbook changed before official export')
    target_excel = O / 'result2.xlsx'
    write_once(target_excel, frozen_excel.read_bytes())
    copy_hash_ok = io.sha(target_excel) == freeze['excel_sha256']
    from openpyxl import load_workbook
    workbook = load_workbook(target_excel, read_only=True, data_only=True)
    try:
        sheet_set_ok = set(workbook.sheetnames) == {'设计参数', '逐镜设计', '镜场设计'}
        eight = list(workbook['镜场设计'].values)
        six = list(workbook['逐镜设计'].values)
        parameters = {row[0]: row[1] for row in list(workbook['设计参数'].values)[1:]}
        parameter_ok = parameters == {'tower_x': design.tower_xy[0], 'tower_y': design.tower_xy[1],
            'width': design.width, 'height': design.height, 'installation_height': design.installation_height, 'n': N}
        count_ok = len(eight) == len(six) == N + 1
        header_ok = list(eight[0]) == OFFICIAL_HEADER and list(six[0]) == ['序号', 'x坐标 (m)', 'y坐标 (m)', '安装高度 (m)', '镜面宽 (m)', '镜面高 (m)']
        field_ok = count_ok and all(
            eight[i + 1] == (*design.tower_xy, i + 1, design.width, design.height, mirror['x'], mirror['y'], design.installation_height)
            and six[i + 1] == (i + 1, mirror['x'], mirror['y'], design.installation_height, design.width, design.height)
            for i, mirror in enumerate(design.mirrors))
        active_ok = workbook.active.title == '镜场设计'
    finally:
        workbook.close()
    readback = {'byte_identical_to_frozen': copy_hash_ok, 'exact_sheet_set': sheet_set_ok,
                'common_parameters_equal': parameter_ok, 'actual_count_equal': count_ok,
                'headers_equal': header_ok, 'all_eight_and_six_fields_equal': field_ok, 'official_sheet_active': active_ok}
    report['official_excel_readback'] = readback
    if not all(readback.values()):
        raise ValueError('Official result2.xlsx independent readback failed: ' + repr(readback))
    table1 = {'header': ['日期', '平均光学效率', '平均余弦效率', '平均阴影遮挡效率', '平均截断效率', '单位面积镜面平均输出热功率 (kW/m²)'],
              'rows': [[str(m + 1) + '月21日', *point[m, :4], point[m, 5]] for m in range(12)]}
    table2 = {'header': ['年平均光学效率', '年平均余弦效率', '年平均阴影遮挡效率', '年平均截断效率', '年平均输出热功率 (MW)', '单位面积镜面年平均输出热功率 (kW/m²)'],
              'rows': [[*point[-1, :4], point[-1, 4] / 1000., point[-1, 5]]]}
    table3 = {'header': ['吸收塔位置坐标 (m)', '定日镜尺寸（宽×高）(m)', '定日镜安装高度 (m)', '定日镜总面数', '定日镜总面积 (m²)'],
              'rows': [['(' + format(design.tower_xy[0], '.17g') + ', ' + format(design.tower_xy[1], '.17g') + ')',
                        format(design.width, '.17g') + ' × ' + format(design.height, '.17g'),
                        design.installation_height, N, design.total_area]]}
    tables = {'table1': table1, 'table2': table2, 'table3': table3}
    csv_receipts = {}
    for name, table in tables.items():
        stream = string_io.StringIO(newline='')
        writer = csv.writer(stream, lineterminator='\n')
        writer.writerow(table['header'])
        writer.writerows(table['rows'])
        path = O / (name + '.csv')
        write_once(path, stream.getvalue().encode('utf-8-sig'))
        with path.open('r', encoding='utf-8-sig', newline='') as handle:
            observed = list(csv.reader(handle))
        expected = [[str(v) for v in table['header']]] + [[str(v) for v in row] for row in table['rows']]
        if observed != expected:
            raise ValueError('CSV roundtrip failed: ' + name)
        csv_receipts[path.name] = {'sha256': io.sha(path), 'rows': len(table['rows']), 'columns': len(table['header']), 'readback_equal': True}
    data = {'version': 'q02-repair-results-v001', 'name': freeze['name'], 'design_key': freeze['design_key'],
            'decision': 'PASS', 'tables': tables, 'internal_power_unit': 'kW', 'table2_total_power_unit': 'MW',
            'actual_design': {'tower_xy': design.tower_xy, 'width_m': design.width, 'height_m': design.height,
                              'installation_height_m': design.installation_height, 'n': N, 'total_area_m2': design.total_area},
            'annual_point_internal_units': point[-1], 'annual_work_indicator_internal_units': U[-1],
            'monthly_and_annual_work_indicator': U, 'rated_power_lower_work_value_kw': lower,
            'sources': {'confirmation_summary_sha256': summary_sha, 'frozen_candidate_sha256': freeze_sha,
                        'independent_point_verification_sha256': io.sha(verification_path),
                        'frozen_excel_sha256': freeze['excel_sha256'], 'result2_excel_sha256': io.sha(target_excel)},
            'csv_receipts': csv_receipts, 'official_excel_readback': readback,
            'evidence_scope': 'Frozen finite search candidate passed prescribed numerical work gates; no global-optimality or physical-certainty claim.'}
    json_path = O / '结果数据.json'
    encoded = (json.dumps(io.convert(data), ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode('utf-8')
    write_once(json_path, encoded)
    if io.load(json_path) != io.convert(data):
        raise ValueError('Result data JSON roundtrip failed')
    report['exports_created'] = True
    report['exports'] = {**csv_receipts, 'result2.xlsx': {'sha256': io.sha(target_excel), 'byte_identical_to_frozen': True},
                         '结果数据.json': {'sha256': io.sha(json_path), 'readback_equal': True}}


def main():
    report = {'version': 'q02-repair-postprocess-v001', 'created': io.now(), 'checks': {},
              'all_pass': False, 'export_allowed': False, 'exports_created': False}
    budget = io.Budget('independent_work_indicator_and_export', 'confirmation', START)
    try:
        run(budget, report)
    except Exception:
        report['execution_error'] = traceback.format_exc()
        report['all_pass'] = False
        report['exports_created'] = False
        report['incomplete'] = True
    finally:
        report['elapsed_seconds_before_save'] = time.perf_counter() - START
        report['budget_used_seconds_before_save'] = budget.used()
        report['postprocess_sha256'] = io.sha(Path(__file__))
        io.save(io.OUT / '独立工作指标重建.json', report)
        budget.finish('FAILED' if report.get('execution_error') else 'COMPLETED')
    print(json.dumps({'all_pass': report['all_pass'], 'export_allowed': report['export_allowed'],
                      'exports_created': report['exports_created'], 'output': str(io.OUT / '独立工作指标重建.json')}, ensure_ascii=False))
    return report


if __name__ == '__main__':
    main()
