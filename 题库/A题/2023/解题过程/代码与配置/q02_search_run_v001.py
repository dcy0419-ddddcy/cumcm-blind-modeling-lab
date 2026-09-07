"""Frozen bounded Q2-A search and holdout confirmation. Run one explicit phase."""
import time
START=time.perf_counter()
from pathlib import Path
import sys,traceback,json,copy,math
HERE=Path(__file__).resolve().parent;sys.path.insert(0,str(HERE))
import q02_search_common_v001 as io
phase=sys.argv[1] if len(sys.argv)>1 else 'invalid'
budget=io.Budget('controlled_'+phase,'confirmation' if phase=='confirm' else 'search',START)
try:
 import numpy as np
 import q02_search_engine_v001 as eng
 import q02_search_summary_tests_v002 as summary_tests
 O=io.OUT;C=O/'candidates';F=eng.f;D=F.design_module;compact=eng.compact
 def initial_specs():
  configs=[((0,0),8,8,6,1.005),((0,-100),8,8,6,1.005),((0,-180),8,8,6,1.005),((80,-160),8,8,6,1.005),((0,-100),7,6,4.5,1.005),((0,-180),7,7,5,1.005),((0,-100),6,6,4,1.005),((0,-180),8,7,5.5,1.005),((80,-160),8,7,5.5,1.005),((0,-180),7.5,7.5,5.5,1.005),((0,-220),8,8,5,1.005),((0,-100),8,8,5,1.04)]
  return [compact.make_ring_spec(t,w,h,z,radial_gap_factor=gap,angular_gap_factor=gap,phase=.5,inner=100.5,spacing_margin=.05,name='C%02d'%(i+1))for i,(t,w,h,z,gap)in enumerate(configs)]
 def rate(spec,kind):
  started=time.perf_counter();folder=C/spec['name'];fd=eng.generate(spec,folder,budget)
  plan=cfg[kind];times,B,n,levels,ns=(plan[k]for k in ('times','B','n','levels','namespace'))
  arrays,recs=eng.evaluate(fd,folder,kind,times,B,n,levels,ns,budget)
  score=eng.quick_score(arrays);score.update(name=fd.name,design_key=fd.key,n=fd.n,area=fd.total_area,spec=spec,score_kind=kind)
  if kind=='fine':
   summary=eng.stored_summary(arrays,folder/kind,'full_comparison')
   score['full_comparison_point']=summary['point'][-1];score['comparison_states']=summary['pooled_state_counts'];score['summary_path']=str(folder/kind/'summary.json')
  score['rate_wall_seconds']=time.perf_counter()-started
  score['evaluation_seconds']=max(score['rate_wall_seconds'],sum(r['prepare_seconds']+r['integral_seconds']+r['save_seconds_before_metadata']for r in recs))
  io.save(folder/kind/'score.json',score);io.event('搜索轨迹.jsonl',score);budget.tick();return io.convert(score)
 def verify_code_bindings(cfg):
  for name,digest in cfg['bindings'].items():
   if io.sha(HERE/name)!=digest:raise RuntimeError('FROZEN_CODE_BINDING_CHANGED: '+name)
 if phase=='search':
  if (O/'最终拟提交候选冻结.json').exists()or(O/'frozen/confirmation/binding.json').exists():raise RuntimeError('CANDIDATE_ALREADY_FROZEN: search cannot change a frozen or confirmed candidate')
  pre=io.load(O/'密集压力与计时-v002.json')
  if not pre['all_pass']:raise RuntimeError('PRECHECKS_NOT_PASSED')
  test=summary_tests.run_checks();io.save(O/'搜索汇总人工测试.json',test)
  if not test['all_pass']:raise RuntimeError('SUMMARY_TEST_FAILURE')
  cfg={'version':'v001','created':io.now(),'before_first_search_score':True,'approved':'C21-C24, D025','initial_specs':initial_specs(),'max_initial':12,'max_local':4,'coarse':{'times':eng.COARSE_TIMES,'B':2,'n':32,'levels':[16,32],'namespace':900,'purpose':'heuristic only, all mirrors at 12 selected times; no rating or deterministic impossibility'},'fine':{'times':eng.TIMES,'B':4,'n':64,'levels':[32,64],'namespace':910,'max_initial_shortlist':2,'coarse_rejection_audit':1,'max_local_shortlist':1},'local_moves':['tower x +30, y +20; dimensions and installation height retained; both ring gaps1.005, phase .5','tower x -30, y -20; dimensions retained; installation height increased .5 capped6; both ring gaps1.005, phase .5','height dimension decreased .5 floored2; installation height6; both ring gaps1.005, phase .5','both ring gaps1.035; phase .25; installation height6'],'selection':'initial two highest q_score among P_score>=62000, else highest P_score; one highest-q coarse exclusion with P<62000, else next excluded ranked; fine among P>=60000 highest q; local starts best fine and chooses top coarse feasible by same rule; final highest fine q among estimated P>=60000, prefer margin if numerical tie','deterministic_elimination':'only explicit geometry/model-domain failure; optical low scores remain heuristic exclusions','random_pairing':'CRN keyed by position_key and standard60 time index and batch; same normalized coordinates/directions for shared keys, not necessarily same physical mirror; different namespaces independent','confirmation':{'B':8,'n':256,'levels':[128,256],'namespace':990,'fixed_sample_no_adaptive_stop':True,'min_reserve_seconds':1200,'U':'max(t7 whole-batch jackknife halfwidth, final-half-prefix difference, final-raw-integral difference) + coordinatewise full-range low-survivor(<100) influence bound','efficiency_abs':.001,'power_relative':.005,'rating':'annual P - U_P >=60000kW and all report metrics pass and no required unresolved state','sample_zero_rule':'no silent zero/one; pooled zeros or unknown are unresolved absent validated independent certificate; no seed retries','low_survival_threshold':100,'interval':'pointwise approximate t7 diagnostic only; U is numerical work indicator, not strict confidence bound or physical uncertainty'},'search_limits':{'tower':'listed initials and four bounded local transformations within fixed circle','dimensions':'listed initial and local values within2..8, width>=height','height':'listed initial/local2..6 with strict z>h/2','layout':'full tower-centered rings, single full azimuth sector, exact chord counts, clipping to fixed field and exclusion; no unrestricted point moves in this run','positive_spacing_search_margin_m':.05,'initial_radial_angular_gap_factors':[1.005,1.04],'coordinate_decimal_places':10},'budget':{'total_seconds':3600,'search_preflight_limit_seconds':2400,'min_confirmation_reserve_seconds':1200,'initial_fine_start_estimate_seconds':400,'minimum_coarse_start_estimate_seconds':120,'freeze_start_estimate_seconds':20,'stop_rule':'per-time/object checkpoint guards; initial coarse preserves initial fine and freeze estimates; skip new stage if extrapolated cost plus freeze estimate exceeds remaining search allocation; estimates are not deterministic runtime bounds'},'bindings':{p.name:io.sha(p)for p in [Path(__file__),HERE/'q02_search_engine_v001.py',HERE/'q02_fast_v002.py',HERE/'q02_compact_v001.py',HERE/'q02_search_summary_v002.py',HERE/'q02_search_common_v001.py']}}
  cp=O/'搜索与确认冻结配置-v001.json'
  if cp.exists():cfg=io.load(cp)
  else:io.save(cp,cfg)
  verify_code_bindings(cfg)
  coarse=[];fine=[];fail=[]
  for spec in cfg['initial_specs']:
   coarse_est=max(120.,max((r['evaluation_seconds']for r in coarse),default=0.)*1.3)
   if not(C/spec['name']/'coarse/score.json').exists()and budget.used()+coarse_est+400+20>=2400:
    io.event('预算停止.jsonl',{'action':'skip_remaining_initial_coarse','next_name':spec['name'],'used_seconds':budget.used(),'coarse_estimate_seconds':coarse_est,'reserved_initial_fine_seconds':400,'reserved_freeze_seconds':20});break
   try:coarse.append(rate(spec,'coarse'))
   except eng.CandidateRejected as ex:
    fail.append({'name':spec['name'],'reason':str(ex),'kind':'input_geometry_or_domain'});io.save(O/'候选失败.json',fail)
  def rank(rows):
   eligible=[r for r in rows if r['P_score_kw']>=62000 and r['unknown']==0]
   return sorted(eligible,key=lambda r:(-r['q_score_kw_m2'],-r['P_score_kw'],r['name']))+sorted([r for r in rows if r not in eligible],key=lambda r:(-r['P_score_kw'],r['name']))
  ranked=rank(coarse);chosen=ranked[:2]
  rejected=[r for r in coarse if r not in chosen]
  if rejected:
   audit=sorted([r for r in rejected if r['P_score_kw']<62000]or rejected,key=lambda r:(-r['q_score_kw_m2'],r['name']))[0];chosen.append(audit)
  io.save(O/'初始粗筛与回查预登记.json',{'coarse':coarse,'selected_full':[r['name']for r in chosen],'audit':chosen[-1]['name']if len(chosen)>2 else None,'exclusions':[{'name':r['name'],'reason':'heuristic score not selected; not proved infeasible'}for r in coarse if r not in chosen]})
  for row in chosen:
   # Conservative per-time extrapolation updated after first complete fine candidate.
   est=400 if not fine else fine[-1].get('evaluation_seconds',400)*1.3
   if not(C/row['name']/'fine/score.json').exists()and budget.used()+est+20>=2400:break
   r=rate(row['spec'],'fine');fine.append(r);io.save(O/'较精比较.json',fine)
  feasible=[r for r in fine if r['P_score_kw']>=60000 and r['unknown']==0]
  if feasible and budget.used()+500<2400:
   base=max(feasible,key=lambda r:r['q_score_kw_m2']);s=base['spec'];tower=s['tower_xy'];w=s['width'];h=s['height'];z=s['installation_height']
   moves=[((tower[0]+30,tower[1]+20),w,h,z,1.005,.5),((tower[0]-30,tower[1]-20),w,h,min(6,z+.5),1.005,.5),(tower,w,max(2,h-.5),6,1.005,.5),(tower,w,h,6,1.035,.25)]
   local_specs=[compact.make_ring_spec(t,w,h,z,radial_gap_factor=gap,angular_gap_factor=gap,phase=ph,inner=100.5,spacing_margin=.05,name='L%02d'%(i+1))for i,(t,w,h,z,gap,ph)in enumerate(moves)]
   io.save(O/'局部调整预登记.json',{'base':base['name'],'specs':local_specs,'before_local_optical_scores':True})
   local=[]
   for sp in local_specs:
    if not(C/sp['name']/'coarse/score.json').exists()and budget.used()+max(120.,max((r['evaluation_seconds']for r in local),default=0.)*1.3)+20>=2400:break
    try:local.append(rate(sp,'coarse'))
    except eng.CandidateRejected as ex:io.event('局部排除.jsonl',{'name':sp['name'],'error':str(ex)})
   if local and budget.used()+max(400.,max((r['evaluation_seconds']for r in fine),default=0.)*1.3)+20<2400:
    r=rate(rank(local)[0]['spec'],'fine');fine.append(r);io.save(O/'较精比较.json',fine)
  feasible=[r for r in fine if r['P_score_kw']>=60000 and r['unknown']==0]
  if not feasible:
   io.save(O/'搜索结束.json',{'status':'NO_ESTIMATED_FEASIBLE_FULL_CANDIDATE','coarse_count':len(coarse),'fine':fine,'not_impossibility_proof':True});budget.finish();sys.exit(0)
  best=max(feasible,key=lambda r:r['q_score_kw_m2'])
  # Freeze actual exported Excel precision BEFORE holdout; preserve stable IDs in separate mapping.
  from openpyxl import Workbook,load_workbook
  fd,read=eng.load_design(C/best['name']);des=fd.to_design();dest=O/'frozen';dest.mkdir(exist_ok=True)
  book=Workbook();ws=book.active;ws.title='设计参数';ws.append(['参数','值','单位']);
  for k,v in [('tower_x',des.tower_xy[0]),('tower_y',des.tower_xy[1]),('width',des.width),('height',des.height),('installation_height',des.installation_height),('n',des.n)]:ws.append([k,v,'m'if k!='n'else'count'])
  ms=book.create_sheet('逐镜设计');ms.append(['序号','x坐标 (m)','y坐标 (m)','安装高度 (m)','镜面宽 (m)','镜面高 (m)'])
  for i,r in enumerate(des.mirrors,1):ms.append([i,r['x'],r['y'],des.installation_height,des.width,des.height])
  path=dest/(best['name']+'-未确认设计.xlsx');book.save(path)
  rb=load_workbook(path,data_only=True,read_only=True);pars={r[0]:r[1]for r in list(rb['设计参数'].values)[1:]};mr=list(rb['逐镜设计'].values)[1:];rb.close()
  if len(mr)!=des.n or pars.get('n')!=des.n:raise RuntimeError('EXCEL_MIRROR_COUNT_MISMATCH')
  expected_parameters={'tower_x':des.tower_xy[0],'tower_y':des.tower_xy[1],'width':des.width,'height':des.height,'installation_height':des.installation_height,'n':des.n}
  if pars!=expected_parameters:raise RuntimeError('EXCEL_COMMON_PARAMETERS_MISMATCH')
  mapped=[]
  for i,row in enumerate(mr):
   if row[0]!=i+1 or row[3:]!=(des.installation_height,des.width,des.height):raise RuntimeError('EXCEL_FIELDS_MISMATCH')
   original=des.mirrors[i];mapped.append({'mirror_id':original['mirror_id'],'position_key':original['position_key'],'x':row[1],'y':row[2],'export_row':i+2})
  outdes=D.Design(des.name,'confirmed-input-v001',(pars['tower_x'],pars['tower_y']),pars['width'],pars['height'],pars['installation_height'],mapped,des.metadata)
  actual=F.freeze_design(outdes);receipt=compact.save_compact(outdes,dest)
  io.save(dest/'身份映射.json',[{'export_index':i+1,'mirror_id':r['mirror_id'],'position_key':r['position_key']}for i,r in enumerate(mapped)])
  io.save(dest/'交付精度读回.json',{'excel':str(path),'sha256':io.sha(path),'n':len(mapped),'actual_design_key':actual.key,'geometry':actual.validation,'max_coordinate_difference_m':float(np.max(np.abs(actual.centers-fd.centers))),'nominal_original_key':fd.key,'actual_total_area':actual.total_area})
  freeze={'name':best['name'],'created':io.now(),'before_confirmation':True,'design_key':actual.key,'bundle_sha256':io.sha(dest/'design.bundle.json'),'excel_sha256':io.sha(path),'best_comparison':best,'all_fine':fine,'confirmation_config':cfg['confirmation'],'search_budget_used_seconds':budget.used(),'comparison_data_used_for_selection':True,'confirmation_data_used_for_selection':False}
  io.save(O/'最终拟提交候选冻结.json',freeze);io.save(O/'搜索结束.json',{'status':'CANDIDATE_FROZEN_NOT_CONFIRMED','coarse_count':len(coarse),'local_count':len(local)if 'local'in locals()else 0,'fine_count':len(fine),'name':best['name']});budget.finish();print(json.dumps({'frozen':best['name'],'fine_P':best['P_score_kw'],'fine_q':best['q_score_kw_m2'],'used':budget.used()},ensure_ascii=False))
 elif phase=='confirm':
  freeze=io.load(O/'最终拟提交候选冻结.json');frozen_cfg=io.load(O/'搜索与确认冻结配置-v001.json')
  verify_code_bindings(frozen_cfg)
  cfg=freeze['confirmation_config']
  if cfg!=frozen_cfg['confirmation']:raise RuntimeError('CONFIRMATION_CONFIG_CHANGED')
  if io.sha(O/'frozen'/(freeze['name']+'-未确认设计.xlsx'))!=freeze['excel_sha256']:raise RuntimeError('FROZEN_EXCEL_CHANGED')
  fd,rr=eng.load_design(O/'frozen')
  if fd.key!=freeze['design_key']or io.sha(O/'frozen/design.bundle.json')!=freeze['bundle_sha256']:raise RuntimeError('FROZEN_INPUT_CHANGED')
  arrays,recs=eng.evaluate(fd,O/'frozen','confirmation',eng.TIMES,cfg['B'],cfg['n'],cfg['levels'],cfg['namespace'],budget)
  ans=eng.stored_summary(arrays,O/'frozen/confirmation','full_confirmation')
  io.save(O/'独立确认结论.json',{'name':fd.name,'design_key':fd.key,'decision':ans['formal_decision'],'annual':ans['point'][-1],'annual_U':ans['conservative_work_indicator'][-1],'all_precision':ans['all_table_items_precision_met'],'rating_lower_kw':ans['rated_power_lower_work_value_kw'],'unresolved':ans['required_table_na_locations'],'low_survival':ans['low_survival'],'numeric_checks':ans['numeric_checks'],'reconstructed_from_saved_chunks':True,'independent_batch_count':cfg['B'],'full_combinations':fd.n*60,'source_samples':fd.n*60*cfg['B']*cfg['n']})
  budget.finish();print(json.dumps(io.load(O/'独立确认结论.json'),ensure_ascii=False))
 else:raise ValueError('phase must be search or confirm')
except SystemExit:
 raise
except BaseException:
 io.event('失败运行.jsonl',{'time':io.now(),'mode':phase,'error':traceback.format_exc(),'used':budget.used()});budget.finish('FAILED');raise



