"""Independent persisted Q2 repair confirmation audit; no rays, search or summary calls.

Run this file only through the root's cumulative computation budget. Inputs are
read-only. Only the shared budget ledger and O/独立重建核验.json are written.
This is file/geometry/arithmetical reconstruction, not physical validation and
not an independent derivation of the reported jackknife/work indicator.
"""
from __future__ import annotations
import time
START = time.perf_counter()
from pathlib import Path
import sys
import math
import hashlib
import traceback

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import q02_repair_common_v001 as io

TIMES = [(m, h) for m in range(1, 13) for h in (9., 10.5, 12., 13.5, 15.)]
LABELS = ['optical', 'cosine', 'shadow_blocking', 'truncation', 'power_kw', 'unit_power_kw_m2']
COUNT_LABELS = ['shadow', 'blocked_after_unshadowed', 'survive', 'capture', 'unknown', 'raw_R']
STATE_LABELS = {0: 'ESTIMATED', 4: 'SAMPLED_ZERO_SURVIVOR', 5: 'SAMPLED_ZERO_CAPTURE',
                6: 'BOUNDARY_UNCERTAIN', 7: 'ZERO_PRIMARY_DENOMINATOR', 8: 'INVALID_STATISTICS'}
WORKER_DEPENDENCIES = {
    'q02_repair_engine_v001.py', 'q02_fast_v002.py', 'q02_repair_common_v001.py',
    'q02_compact_v001.py', 'q02_hex_bands_v001.py', 'q01_core_v003.py',
    'q02_eval_v003.py', 'q02_design_v002.py', 'q02_repair_parallel_v001.py',
}


def audit_parallel_schedule(budget, report, check, cfg, freeze, binding, source):
    """Read real scheduler receipts; never import scheduler or inspect processes.

    run/job receipts have no separate design_key or chunk digest. Their source
    identities are checked against the global binding, then tasks are connected
    to the same directory's design-bound time metadata. The main audit separately
    verifies those metadata/chunk hashes. Failed historical attempts are retained;
    completed, same-source task receipts can cover resumed time chunks.
    """
    execution = cfg.get('execution', {})
    expected_scheduler = 'q02_repair_parallel_v001.py'
    scheduler = execution.get('scheduler')
    configured = cfg.get('bindings', {})
    dependencies = binding.get('worker_dependency_sha256')
    valid_registry = isinstance(dependencies, dict) and set(dependencies) == WORKER_DEPENDENCIES
    check('scheduling', 'exact_nine_worker_dependency_registry', valid_registry,
          {'expected': sorted(WORKER_DEPENDENCIES),
           'actual': sorted(dependencies) if isinstance(dependencies, dict) else None})
    if not isinstance(dependencies, dict):
        dependencies = {}
    check('scheduling', 'scheduler_matches_frozen_configuration',
          scheduler == expected_scheduler and expected_scheduler in configured
          and binding.get('scheduler_sha') == configured.get(expected_scheduler),
          {'configured_scheduler': scheduler, 'binding_sha256': binding.get('scheduler_sha'),
           'configuration_sha256': configured.get(expected_scheduler)})
    check('scheduling', 'four_workers_match_execution',
          type(binding.get('workers')) is int and type(execution.get('workers')) is int
          and binding['workers'] == execution['workers'] == 4)
    check('scheduling', 'reserved_tail_matches_execution',
          binding.get('reserve_after_seconds') == execution.get('reserve_after_seconds'))
    check('scheduling', 'bound_compact_input',
          binding.get('compact_bundle_sha256') == freeze.get('bundle_sha256'))
    actual_hashes = {}
    for name in sorted(WORKER_DEPENDENCIES):
        budget.guard(5)
        path = HERE / name
        actual = io.sha(path) if path.is_file() else None
        actual_hashes[name] = actual
        check('scheduling', 'actual_worker_source_' + name,
              actual is not None and actual == dependencies.get(name),
              {'file_exists': path.is_file(), 'actual_sha256': actual,
               'worker_bound_sha256': dependencies.get(name)})
        if name in configured:
            check('scheduling', 'worker_source_matches_cfg_' + name,
                  dependencies.get(name) == configured[name])
    for name, field in [('q02_repair_engine_v001.py', 'code_sha'),
                        ('q02_fast_v002.py', 'fast_sha'),
                        ('q02_repair_common_v001.py', 'common_sha'),
                        ('q02_compact_v001.py', 'compact_sha'),
                        (expected_scheduler, 'scheduler_sha')]:
        check('scheduling', 'worker_and_primary_binding_' + name,
              name in configured and dependencies.get(name) == binding.get(field) == configured.get(name))

    audit_dir = source / 'parallel-audit'
    run_paths = sorted((p / 'run.json' for p in audit_dir.glob('run-*') if p.is_dir()),
                       key=lambda p: p.parent.name)
    check('scheduling', 'parallel_run_receipts_exist', bool(run_paths))
    coverage = set()
    histories = []
    current = None
    current_tasks = []
    current_final_ok = False
    metadata_cache = {}

    def load_receipt(path):
        try:
            value = io.load(path)
            return value if isinstance(value, dict) else None
        except (OSError, ValueError):
            return None

    for ordinal, path in enumerate(run_paths):
        budget.guard(5)
        is_current = ordinal == len(run_paths) - 1
        run = load_receipt(path)
        if run is None:
            histories.append({'run': str(path.relative_to(io.ROOT)), 'readable': False})
            if is_current:
                check('scheduling', 'latest_run_receipt_readable', False)
            continue
        matching = (valid_registry and run.get('worker_source_sha256') == dependencies
                    and run.get('version') == binding.get('scheduler_version')
                    and run.get('workers') == execution.get('workers') == 4)
        init_paths = sorted(path.parent.glob('init-*.json'))
        task_paths = sorted(path.parent.glob('job-*.json'))
        initializers = [load_receipt(p) for p in init_paths]
        tasks = [load_receipt(p) for p in task_paths]
        good_initializers = {r['pid']: r for r in initializers
                             if isinstance(r, dict) and type(r.get('pid')) is int
                             and r.get('error') is None and r.get('source_hashes') == dependencies}
        valid_tasks = []
        for task_path, task in zip(task_paths, tasks):
            if not isinstance(task, dict):
                continue
            ti = task.get('local_time_index')
            if type(ti) is not int or not 0 <= ti < 60:
                continue
            if ti not in metadata_cache:
                metadata_cache[ti] = load_receipt(source / ('time-%02d.json' % ti))
            meta = metadata_cache[ti]
            valid = (matching and task_path.name == 'job-%02d.json' % ti
                     and task.get('standard_time') == io.convert(TIMES[ti])
                     and task.get('completed') is True and task.get('error') is None
                     and task.get('pid') in good_initializers and isinstance(meta, dict)
                     and (meta.get('month'), meta.get('hour')) == TIMES[ti]
                     and meta.get('design_key') == binding.get('design_key')
                     and meta.get('combinations') == binding.get('n_mirrors')
                     and meta.get('unique_source_samples') == binding.get('B') * binding.get('n') * binding.get('n_mirrors'))
            if valid:
                valid_tasks.append(ti)
                coverage.add(ti)
        histories.append({'run': str(path.relative_to(io.ROOT)), 'sha256': io.sha(path),
                          'readable': True, 'complete': run.get('complete'),
                          'source_identity_matches': matching, 'same_source_committed_tasks': valid_tasks,
                          'historical_error': run.get('error')})
        if not is_current:
            continue
        current, current_tasks = run, valid_tasks
        check('scheduling', 'latest_run_source_identity', matching)
        current_final_ok = (matching and run.get('complete') is True and isinstance(run.get('finished'), str)
              and bool(run['finished']) and not run.get('error')
              and run.get('terminated_unfinished_worker_cpu_may_be_unavailable') is False)
        check('scheduling', 'latest_run_final_completion', current_final_ok)
        init_pids = [r.get('pid') for r in initializers if isinstance(r, dict)]
        aggregate_init = run.get('initializer_receipts')
        aggregate_init_valid = (isinstance(aggregate_init, list)
                                and all(isinstance(r, dict) and type(r.get('pid')) is int
                                        for r in aggregate_init))
        check('scheduling', 'latest_initializer_receipts',
              bool(initializers) and len(good_initializers) == len(initializers)
              and len(init_pids) == len(set(init_pids))
              and run.get('workers_initialized') == len(initializers)
              and aggregate_init_valid
              and sorted(aggregate_init, key=lambda r: r['pid'])
                  == sorted(initializers, key=lambda r: r['pid']))
        aggregate = run.get('worker_receipts')
        aggregate_valid = (isinstance(aggregate, list) and all(isinstance(r, dict) for r in aggregate)
                           and all(type(r.get('local_time_index')) is int for r in aggregate))
        check('scheduling', 'latest_task_receipts_match_disk',
              aggregate_valid and all(isinstance(r, dict) and type(r.get('local_time_index')) is int
                                      for r in tasks)
              and run.get('task_receipts') == len(tasks)
              and sorted(aggregate, key=lambda r: r['local_time_index'])
                  == sorted(tasks, key=lambda r: r['local_time_index']))
        completed = run.get('completed_local_indices')
        valid_completed = (isinstance(completed, list) and all(type(i) is int and 0 <= i < 60 for i in completed))
        check('scheduling', 'latest_successful_submitted_times',
              valid_completed and len(completed) == len(set(completed))
              and len(valid_tasks) == len(tasks) and len(valid_tasks) == len(set(valid_tasks))
              and sorted(completed) == sorted(valid_tasks),
              {'completed_local_indices': completed, 'validated_task_indices': valid_tasks})

    check('scheduling', 'all_60_times_have_same_source_completed_task_receipts', coverage == set(range(60)),
          {'covered_time_indices': sorted(coverage), 'missing_time_indices': sorted(set(range(60)) - coverage)})
    report['parallel_schedule_audit'] = {
        'execution_configuration': execution, 'worker_dependency_actual_sha256': actual_hashes,
        'attempts': histories, 'latest_run': str(run_paths[-1].relative_to(io.ROOT)) if run_paths else None,
        'latest_run_valid_task_count': len(current_tasks), 'all_attempts_valid_time_count': len(coverage),
        'latest_run_itself_covers_all_60': set(current_tasks) == set(range(60)),
        'same_source_failed_history_may_supply_committed_resume_chunks': True,
        'identity_chain': 'run/initializer source hashes -> same-directory confirmation binding -> per-time design metadata; chunk hashes verified separately below',
        'receipt_exit_path_evidence': current_final_ok,
        'exit_evidence_scope': 'Under the bound scheduler, complete is saved after child-exit waiting and Pool.join; this is saved control-flow evidence only.',
        'operating_system_current_worker_state_checked': False,
        'no_claim_of_current_worker_absence': True,
        'worker_receipts_do_not_contain_separate_design_key_or_chunk_digest': True,
        'wall_accounting_scope': binding.get('budget_accounting'),
        'worker_cpu_receipts_are_lower_bounds_not_additional_wall_time': True,
    }


def audit(budget, report):
    import numpy as np
    import q02_fast_v002 as fast
    import q02_compact_v001 as compact

    checks = report['checks']
    def check(group, name, passed, evidence=None):
        checks.append({'group': group, 'name': name, 'passed': bool(passed), 'evidence': io.convert(evidence)})
        return bool(passed)

    def compare(group, name, actual, expected, atol=2e-12, rtol=2e-12, required=False):
        a = np.asarray(actual, dtype=float)
        b = np.asarray(expected, dtype=float)
        if a.shape != b.shape:
            return check(group, name, False, {'actual_shape': a.shape, 'expected_shape': b.shape})
        af, bf = np.isfinite(a), np.isfinite(b)
        finite = af & bf
        same_na = np.array_equal(af, bf)
        err = float(np.max(np.abs(a[finite] - b[finite]))) if np.any(finite) else None
        close = bool(np.all(np.abs(a[finite] - b[finite]) <= atol + rtol * np.abs(b[finite])))
        passed = same_na and close and (not required or bool(af.all()))
        return check(group, name, passed, {'max_abs_error': err, 'atol': atol, 'rtol': rtol,
            'na_pattern_equal': same_na, 'na_count_actual': int((~af).sum()),
            'finite_compared': int(finite.sum()), 'required_all_defined': required})

    O = io.OUT
    frozen_dir = O / 'frozen'
    source = frozen_dir / 'confirmation'
    freeze = io.load(O / '最终拟提交候选冻结.json')
    cfg = io.load(O / freeze['config_file'])
    binding = io.load(source / 'binding.json')
    summary = io.load(source / 'summary.json')
    checkpoint = io.load(source / 'checkpoint.json')
    report['input_paths'] = {'frozen_candidate': str(O / '最终拟提交候选冻结.json'),
        'configuration': str(O / freeze['config_file']), 'confirmation': str(source)}
    report['source_hashes'] = {p.name: io.sha(p) for p in
        (O / '最终拟提交候选冻结.json', O / freeze['config_file'],
         source / 'binding.json', source / 'summary.json', source / 'checkpoint.json')}
    for name, digest in cfg['bindings'].items():
        check('binding', 'frozen_code_' + name, io.sha(HERE / name) == digest)
    engine_name = next(name for name in cfg['bindings'] if name.startswith('q02_repair_engine_'))
    check('binding', 'engine_hash', binding.get('code_sha') == io.sha(HERE / engine_name))
    check('binding', 'fast_hash', binding.get('fast_sha') == io.sha(HERE / 'q02_fast_v002.py'))
    check('binding', 'confirmation_config', freeze.get('confirmation_config') == cfg.get('confirmation'))
    audit_parallel_schedule(budget, report, check, cfg, freeze, binding, source)
    plan = cfg['confirmation']
    check('binding', 'confirmation_sample_plan',
          (binding.get('B'), binding.get('n'), binding.get('levels')) == (8, 256, [128, 256])
          and (plan.get('B'), plan.get('n'), plan.get('levels')) == (8, 256, [128, 256]))
    confirm_namespace = freeze['confirmation_config']['namespace']
    root_seed = cfg['root_seed']
    search_namespaces = cfg['search_namespaces']
    check('binding', 'search_namespace_registry', isinstance(search_namespaces, list)
          and bool(search_namespaces) and all(type(v) is int and v >= 0 for v in search_namespaces)
          and len(set(search_namespaces)) == len(search_namespaces))
    check('binding', 'holdout_namespace', type(confirm_namespace) is int and confirm_namespace >= 0
          and binding.get('namespace') == plan.get('namespace') == confirm_namespace
          and confirm_namespace not in search_namespaces)
    # Repair runner records the actual selected fine score and ranking basis;
    # it does not emit the older comparison_data_used_for_selection flag.
    comparison = freeze.get('best_comparison')
    ranking_basis = freeze.get('ranking_basis')
    selection_evidence = (isinstance(comparison, dict) and comparison.get('name') == freeze.get('name')
        and comparison.get('score_kind') == 'fine' and comparison.get('actual_full60') is True
        and comparison.get('fine_gate') is True and isinstance(ranking_basis, str) and bool(ranking_basis.strip()))
    check('binding', 'selection_before_confirmation', freeze.get('before_confirmation') is True
          and freeze.get('confirmation_data_used_for_selection') is False and selection_evidence,
          {'evidence_fields': ['best_comparison', 'ranking_basis', 'confirmation_data_used_for_selection'],
           'legacy_comparison_flag_present': 'comparison_data_used_for_selection' in freeze})
    report['interface_adaptation'] = {
        'reason': 'Frozen repair runner records best_comparison and ranking_basis instead of the older comparison-data flag.',
        'actual_ranking_basis': ranking_basis, 'actual_selected_comparison_name': comparison.get('name') if isinstance(comparison, dict) else None,
        'missing_legacy_flag_inferred_or_written': False}
    check('binding', 'root_seed', type(root_seed) is int and root_seed >= 0 and binding.get('root_seed') == root_seed)
    check('coverage', 'standard_60_binding_times', binding.get('times') == io.convert(TIMES))
    excel = (O / freeze['excel_file']).resolve()
    if excel.parent != frozen_dir.resolve() or excel.suffix.lower() != '.xlsx':
        raise ValueError('Frozen Excel must remain inside this repair output frozen directory')
    check('binding', 'frozen_excel_sha', io.sha(excel) == freeze['excel_sha256'])
    design, receipt = compact.read_compact(frozen_dir / 'design.bundle.json', freeze['bundle_sha256'])
    fd = fast.freeze_design(design)
    N, B, levels = fd.n, 8, np.asarray([128, 256], dtype=np.int64)
    area = float(N * fd.width * fd.height)
    report['compact_readback'] = {k: receipt[k] for k in
        ('bundle_sha256', 'metadata_sha256', 'mirrors_sha256', 'canonical_design_sha256',
         'all_bound_fields_equal', 'actual_n', 'geometry_validation_performed')}
    from openpyxl import load_workbook
    book = load_workbook(excel,read_only=True,data_only=True)
    excel_rows = list(book['逐镜设计'].values)[1:]
    check('binding','excel_rows_independently_read', len(excel_rows)==fd.n)
    check('binding','required_eight_column_sheet','镜场设计' in book.sheetnames)
    pars = {row[0]:row[1] for row in list(book['设计参数'].values)[1:]}
    check('binding','global_parameter_sheet',pars=={'tower_x':fd.tower_xy[0],'tower_y':fd.tower_xy[1],'width':fd.width,'height':fd.height,'installation_height':fd.installation_height,'n':fd.n})
    if '镜场设计' in book.sheetnames:
        all_eight = list(book['镜场设计'].values)
        check('binding','eight_column_template_header',list(all_eight[0])==['吸收塔x坐标 (m)','吸收塔y坐标 (m)','定日镜序号','定日镜宽度 (m)','定日镜高度 (m)','定日镜x坐标 (m)','定日镜y坐标 (m)','定日镜z坐标 (m)'])
        eight = all_eight[1:]
        check('binding','eight_column_field_mapping', len(eight)==fd.n and all(row==(*fd.tower_xy,i+1,fd.width,fd.height,float(fd.centers[i,0]),float(fd.centers[i,1]),fd.installation_height) for i,row in enumerate(eight)))
    check('binding','excel_actual_geometry', len(excel_rows)==fd.n and all(row==(i+1,float(fd.centers[i,0]),float(fd.centers[i,1]),fd.installation_height,fd.width,fd.height) for i,row in enumerate(excel_rows)))
    book.close()
    check('binding', 'actual_frozen_design_key', fd.key == freeze['design_key'] == binding['design_key'])
    check('coverage', 'actual_count_and_name', binding.get('n_mirrors') == N
          and binding.get('name') == fd.name == freeze['name'])
    check('coverage', 'unique_internal_ids_and_keys', len(set(fd.mirror_ids)) == N
          and len(set(fd.position_keys)) == N and N > 0)
    mapping = io.load(frozen_dir / '身份映射.json')
    expected_mapping = [{'export_index': i + 1, 'mirror_id': ident, 'position_key': key}
                        for i, (ident, key) in enumerate(zip(fd.mirror_ids, fd.position_keys))]
    check('binding', 'frozen_identity_mapping', mapping == expected_mapping)
    digests = [hashlib.sha256(key.encode()).digest()[:16] for key in fd.position_keys]
    check('binding', 'unique_position_seed_digests', len(set(digests)) == N)
    seed_examples = []
    for ti, key, batch in ((0, fd.position_keys[0], 0), (59, fd.position_keys[-1], 7)):
        words = fast.reference.seed_words(root_seed, confirm_namespace, ti, key, batch)
        expected = [root_seed, confirm_namespace, ti] + [int.from_bytes(hashlib.sha256(key.encode()).digest()[j:j+4], 'little') for j in range(0, 16, 4)] + [batch]
        check('binding', 'seed_word_encoding_%d_%d' % (ti, batch), words == expected)
        seed_examples.append({'time_index': ti, 'batch': batch, 'words': words})
    report['randomness_audit'] = {'root_seed': root_seed, 'confirmation_namespace': confirm_namespace,
        'search_namespaces': search_namespaces, 'examples': seed_examples, 'all_position_digests_unique': len(set(digests)) == N,
        'claim': 'Distinct deterministic namespace/position/time/batch labels; no generated samples or statistical independence theorem.'}

    # Independently check approved geometry directly from reconstructed centers.
    xy = np.asarray(fd.centers[:, :2])
    tower = np.asarray(fd.tower_xy)
    field_radius = np.hypot(xy[:, 0], xy[:, 1])
    tower_radius = np.hypot(xy[:, 0] - tower[0], xy[:, 1] - tower[1])
    dims_ok = 2 <= fd.width <= 8 and 2 <= fd.height <= 8 and 2 <= fd.installation_height <= 6
    check('geometry', 'fixed_constraints', dims_ok and fd.width >= fd.height
          and fd.installation_height > fd.height / 2 and math.hypot(*fd.tower_xy) <= 350
          and bool(np.all(field_radius <= 350)) and bool(np.all(tower_radius >= 100)),
          {'actual_n': N, 'width': fd.width, 'height': fd.height, 'installation_height': fd.installation_height,
           'max_center_radius_m': float(field_radius.max()), 'min_tower_distance_m': float(tower_radius.min()),
           'ground_gap_m': fd.installation_height - fd.height / 2,
           'scope': 'Center boundaries and approved width>=height search scope, no full-panel boundary substitution.'})
    min_distance = math.inf
    pairs = 0
    for start in range(0, N, 128):
        budget.guard(5)
        rows = np.arange(start, min(start + 128, N))
        dx = xy[rows, 0, None] - xy[None, :, 0]
        dy = xy[rows, 1, None] - xy[None, :, 1]
        distance = np.hypot(dx, dy)
        keep = rows[:, None] < np.arange(N)[None, :]
        if np.any(keep):
            min_distance = min(min_distance, float(distance[keep].min()))
            pairs += int(keep.sum())
    check('geometry', 'independent_all_pairs_spacing', pairs == N * (N - 1) // 2
          and min_distance >= fd.width + 5, {'checked_pairs': pairs,
              'minimum_distance_m': min_distance if math.isfinite(min_distance) else None,
              'required_distance_m': fd.width + 5})
    compare('geometry', 'uniform_actual_areas', fd.areas, np.full(N, fd.width * fd.height), atol=0, rtol=0, required=True)
    compare('geometry', 'total_area', [fd.total_area], [area], atol=1e-8, required=True)

    expected_files = {'time-%02d.npz' % i for i in range(60)}
    check('coverage', 'exact_time_chunk_files', {p.name for p in source.glob('time-*.npz')} == expected_files)
    check('coverage', 'complete_checkpoint', checkpoint.get('complete') is True
          and checkpoint.get('completed_times') == checkpoint.get('expected_times') == 60
          and len(checkpoint.get('records', [])) == 60)
    check('coverage', 'summary_population', summary.get('coverage') == {
        'levels_n_per_batch': [128, 256], 'levels': 2, 'batches': 8, 'times': 60, 'objects': N,
        'final_samples_per_object_time': 2048, 'complete_time_population': True})
    check('coverage', 'summary_mode_is_confirmation', summary.get('kind') == 'Q02_FULL_CONFIRMATION_SUMMARY'
          and summary.get('heuristic_only') is False and summary.get('formal_decision') in {'PASS', 'FAIL', 'UNRESOLVED'})
    check('binding', 'summary_declared_targets', summary.get('targets') == {
        'efficiency_absolute': plan['efficiency_abs'], 'power_relative': plan['power_relative'], 'rated_power_kw': 60000.0})
    check('coverage', 'required_confirmation_fields_present',
          {'formal_ready', 'required_table_na_locations', 'diagnostic_na_locations', 'conservative_work_indicator',
           'all_table_items_precision_met', 'rated_power_lower_work_value_kw'} <= set(summary))
    check('coverage', 'summary_labels', summary.get('labels') == LABELS and summary.get('count_labels') == COUNT_LABELS)
    check('coverage', 'summary_times_and_ids', summary.get('months') == [m for m, h in TIMES]
          and summary.get('hours') == [h for m, h in TIMES] and summary.get('object_ids') == list(fd.mirror_ids))
    check('state', 'no_unhandled_true_zero_certificates', np.asarray(summary.get('certificate_codes')).shape == (60, N)
          and bool(np.all(np.asarray(summary.get('certificate_codes')) == 0)),
          'Current engine supplies no certificates; any added true-zero channel requires separate independent proof audit.')
    temporal = np.full((60, 6), np.nan)
    rebuilt_states = np.full((60, N), 8, dtype=np.int8)
    total_defined = np.zeros((60, N), dtype=bool)
    component_defined = np.zeros((60, N, 4), dtype=bool)
    domain_records, chunk_records, unresolved = [], [], []
    for ti, (month, hour) in enumerate(TIMES):
        budget.guard(5)
        path = source / ('time-%02d.npz' % ti)
        meta_path = path.with_suffix('.json')
        meta = io.load(meta_path)
        digest = io.sha(path)
        check('binding', 'chunk_hash_%02d' % ti, digest == meta.get('sha256'))
        check('binding', 'chunk_metadata_%02d' % ti,
              (meta.get('month'), meta.get('hour'), meta.get('design_key'), meta.get('combinations'), meta.get('unique_source_samples'))
              == (month, hour, fd.key, N, 2048 * N))
        check('binding', 'checkpoint_record_%02d' % ti,
              ti < len(checkpoint.get('records', [])) and checkpoint['records'][ti] == meta)
        scene = fast.scene_at(fd, month, hour)
        check('domain', 'scene_binding_%02d' % ti, scene.key == meta.get('scene_key'))
        domain_records.append({'time_index': ti, 'month': month, 'hour': hour,
                               'scene_key': scene.key, 'domain': scene.domain})
        with np.load(path, allow_pickle=False) as src:
            needed = {'sums', 'cross', 'counts', 'unknown_weights', 'known_capture_weights',
                      'cosine', 'tau', 'dni', 'areas', 'candidate_counts'}
            if not needed <= set(src.files):
                raise ValueError('Missing required arrays in ' + path.name)
            s, ct, uw = src['sums'], src['counts'], src['unknown_weights']
            cross, known, candidates = src['cross'], src['known_capture_weights'], src['candidate_counts']
            co, ta, dni, areas = src['cosine'], src['tau'], src['dni'], src['areas']
            shapes_ok = (s.shape == (2, B, N, 3) and ct.shape == (2, B, N, 6)
                and uw.shape == known.shape == (2, B, N) and cross.shape == (2, B, N, 3, 3)
                and co.shape == ta.shape == areas.shape == (N,) and dni.shape == () and candidates.shape == (N, 2))
            if not check('statistics', 'chunk_shapes_%02d' % ti, shapes_ok):
                continue
            compare('domain', 'cosine_%02d' % ti, co, scene.cosine, atol=1e-13, rtol=1e-13, required=True)
            compare('domain', 'atmosphere_%02d' % ti, ta, scene.tau, atol=1e-13, rtol=1e-13, required=True)
            compare('domain', 'dni_%02d' % ti, dni, scene.dni, atol=1e-13, rtol=1e-13, required=True)
            compare('geometry', 'chunk_areas_%02d' % ti, areas, fd.areas, atol=0, rtol=0, required=True)
            ng = levels[:, None, None]
            finite_s = np.all(np.isfinite(s), axis=-1)
            bad_s = (~finite_s) | np.any(s < 0, axis=-1) | (s[..., 1] > s[..., 0] + 1e-10) | (s[..., 2] > s[..., 1] + 1e-10)
            bad_c = (~np.all(np.isfinite(ct), axis=-1)) | np.any(ct < 0, axis=-1) | np.any(ct > ng[..., None], axis=-1)
            bad_c |= np.any(ct != np.rint(ct), axis=-1) | (ct[..., 0] + ct[..., 1] + ct[..., 2] != ng)
            bad_c |= (ct[..., 3] > ct[..., 2]) | (ct[..., 3] > ct[..., 5])
            bad_w = (~np.isfinite(uw)) | (uw < 0) | (uw > s[..., 0] + 1e-10)
            bad_w |= ((ct[..., 4] == 0) & (uw > 0)) | ((ct[..., 4] > 0) & (uw <= 0))
            bad_w |= ((ct[..., 2] == 0) & (np.abs(s[..., 1]) > 1e-10)) | ((ct[..., 2] > 0) & (s[..., 1] <= 0))
            bad_w |= ((ct[..., 3] == 0) & (np.abs(s[..., 2]) > 1e-10)) | ((ct[..., 3] > 0) & (s[..., 2] <= 0))
            nested = np.any(np.diff(s, axis=0) < -1e-10, axis=-1) | np.any(np.diff(ct, axis=0) < 0, axis=-1) | (np.diff(uw, axis=0) < -1e-10)
            nested_bad = np.any(nested, axis=0)
            check('statistics', 'energies_counts_and_prefixes_%02d' % ti,
                  not np.any(bad_s | bad_c | bad_w) and not np.any(nested_bad),
                  {'bad_energy_rows': int(bad_s.sum()), 'bad_count_rows': int(bad_c.sum()),
                   'bad_weight_rows': int(bad_w.sum()), 'bad_prefix_objects_by_batch': int(nested_bad.sum())})
            known_ok = np.all(np.isfinite(known)) and np.all(known >= 0) and np.all(known <= s[..., 2] + 1e-10)
            known_ok = bool(known_ok and np.all(np.abs(known[ct[..., 4] == 0] - s[..., 2][ct[..., 4] == 0]) <= 1e-10))
            diag = np.diagonal(cross, axis1=-2, axis2=-1)
            cross_ok = bool(np.all(np.isfinite(cross)) and np.all(cross >= 0)
                and np.all(np.abs(cross - cross.swapaxes(-1, -2)) <= 1e-10)
                and np.all(diag * ng[..., None] + 1e-8 >= s * s))
            check('statistics', 'known_capture_and_cross_%02d' % ti, known_ok and cross_ok)
            candidate_ok = bool(np.all(np.isfinite(candidates)) and np.all(candidates == np.rint(candidates))
                                and np.all(candidates >= 0) and np.all(candidates <= N - 1))
            check('statistics', 'candidate_counts_%02d' % ti, candidate_ok)
            compare('binding', 'candidate_metadata_mean_%02d' % ti, candidates.mean(axis=0), meta.get('candidate_mean'), required=True)
            compare('binding', 'candidate_metadata_max_%02d' % ti, candidates.max(axis=0), meta.get('candidate_max'), atol=0, rtol=0, required=True)
            pooled, counts, unknown = s[-1].sum(axis=0), ct[-1].sum(axis=0), uw[-1].sum(axis=0)
            a, b, capture = pooled[:, 0], pooled[:, 1], pooled[:, 2]
            invalid = np.any((bad_s | bad_c | bad_w)[-1] | nested_bad, axis=0)
            state = np.zeros(N, dtype=np.int8)
            state[invalid] = 8
            state[~invalid & (a <= 0)] = 7
            eligible = ~invalid & (a > 0)
            uncertain = (counts[:, 4] > 0) | (unknown > 0)
            state[eligible & uncertain] = 6
            known_domain = eligible & ~uncertain
            state[known_domain & (b <= 0)] = 4
            state[known_domain & (b > 0) & (capture <= 0)] = 5
            est = known_domain & (b > 0) & (capture > 0)
            metrics = np.full((N, 4), np.nan)
            metrics[:, 1] = co
            sb_defined = known_domain & (b > 0)
            metrics[sb_defined, 2] = b[sb_defined] / a[sb_defined]
            metrics[est, 3] = capture[est] / b[est]
            metrics[est, 0] = .92 * co[est] * ta[est] * capture[est] / a[est]
            powers = np.full(N, np.nan)
            powers[est] = float(dni) * areas[est] * metrics[est, 0]
            # Ordinary means/sums preserve every actual mirror, propagating NA.
            temporal[ti, :4] = np.mean(metrics, axis=0)
            temporal[ti, 4] = np.sum(powers)
            temporal[ti, 5] = temporal[ti, 4] / area
            rebuilt_states[ti] = state
            total_defined[ti] = est
            component_defined[ti] = np.isfinite(metrics)
            if np.any(est):
                product = .92 * co[est] * ta[est] * metrics[est, 2] * metrics[est, 3]
                compare('statistics', 'five_factors_%02d' % ti, product, metrics[est, 0], required=True)
                check('statistics', 'efficiency_range_%02d' % ti, bool(np.all(metrics[est] >= 0) and np.all(metrics[est] <= 1 + 1e-10)))
            for code in (4, 5, 6, 7, 8):
                hit = np.flatnonzero(state == code)
                if len(hit):
                    unresolved.append({'time_index': ti, 'state': STATE_LABELS[code], 'count': len(hit),
                        'first_20_ids': [fd.mirror_ids[i] for i in hit[:20]]})
            chunk_records.append({'time_index': ti, 'npz_sha256': digest, 'metadata_sha256': io.sha(meta_path),
                                  'pooled_counts': [int(v) for v in counts.sum(axis=0)]})
        if ti % 5 == 4:
            budget.tick()
    rows = np.vstack([np.mean(temporal.reshape(12, 5, 6), axis=1), np.mean(temporal, axis=0)[None, :]])
    expected_rows = [f'{m:02d}-21' for m in range(1, 13)] + ['annual']
    check('coverage', 'summary_row_labels', summary.get('rows') == expected_rows)
    for ri, row_name in enumerate(expected_rows):
        for col, label in enumerate(LABELS):
            compare('rebuild', row_name + '_' + label, [rows[ri, col]], [summary['point'][ri][col]],
                    atol=1e-7 if col == 4 else 2e-12, rtol=2e-12, required=True)
    compare('rebuild', 'temporal_point', temporal, summary.get('temporal_point'), atol=1e-7, required=True)
    compare('rebuild', 'formal_h02_rows', rows, summary.get('formal_h02_rows'), atol=1e-7, required=True)
    compare('state', 'pooled_state_codes', rebuilt_states, summary.get('pooled_state_codes'), atol=0, rtol=0, required=True)
    compare('state', 'pooled_total_defined', total_defined, summary.get('pooled_total_defined'), atol=0, rtol=0, required=True)
    compare('state', 'pooled_component_defined', component_defined, summary.get('pooled_component_defined'), atol=0, rtol=0, required=True)
    compare('rebuild', 'power_area_identity_all_rows', rows[:, 5] * area, rows[:, 4], atol=1e-7, required=True)
    check('state', 'no_unresolved_objects', not unresolved, {'object_time_count': sum(v['count'] for v in unresolved)})
    check('state', 'no_reported_unresolved_table_or_diagnostic',
          not summary.get('required_table_na_locations') and not summary.get('diagnostic_na_locations')
          and summary.get('formal_decision') != 'UNRESOLVED')
    report.update({'reconstructed_point': rows, 'reconstructed_temporal_point': temporal,
        'labels': LABELS, 'rows': expected_rows, 'actual_n': N, 'actual_total_area_m2': area,
        'covered_combinations': 60 * N, 'unique_confirmatory_source_samples': 60 * N * 2048,
        'source_sample_count_excludes_prefixes_and_rebuild': True,
        'unresolved': unresolved, 'reconstructed_domains_60': domain_records, 'chunks': chunk_records,
        'reported_formal_decision': summary.get('formal_decision'),
        'not_recomputed': ['jackknife uncertainty', 'work indicator and low-survival diameter', 'physical-model discrepancy'],
        'no_new_rays': True, 'no_new_search': True, 'numerical_reconstruction_independent_of_summary_implementation': True})


def main():
    report = {'version': 'q02-repair-verify-v001', 'created': io.now(), 'checks': [],
              'scope': 'Persisted file binding, fixed constraints, geometry domain and point-statistic arithmetic; not physical validation.',
              'all_pass': False}
    budget = io.Budget('independent_rebuild', 'confirmation', START)
    try:
        audit(budget, report)
        report['all_pass'] = bool(report['checks']) and all(c['passed'] for c in report['checks'])
        report['file_and_binding_pass'] = all(c['passed'] for c in report['checks'] if c['group'] in {'binding', 'coverage', 'scheduling'})
        report['scheduling_binding_pass'] = all(c['passed'] for c in report['checks'] if c['group'] == 'scheduling')
        report['numeric_rebuild_pass'] = all(c['passed'] for c in report['checks'] if c['group'] in {'statistics', 'rebuild', 'state'})
        report['geometry_domain_pass'] = all(c['passed'] for c in report['checks'] if c['group'] in {'geometry', 'domain'})
        report['failed_check_count'] = sum(not c['passed'] for c in report['checks'])
    except Exception:
        report['execution_error'] = traceback.format_exc()
        report['all_pass'] = False
        report['incomplete'] = True
    finally:
        report['elapsed_seconds_before_save'] = time.perf_counter() - START
        report['budget_used_seconds_before_save'] = budget.used()
        report['verifier_sha256'] = io.sha(Path(__file__))
        io.save(io.OUT / '独立重建核验.json', report)
        budget.finish('COMPLETED' if not report.get('execution_error') else 'FAILED')
    print(__import__('json').dumps({'all_pass': report['all_pass'],
          'failed_checks': report.get('failed_check_count'), 'incomplete': report.get('incomplete', False),
          'output': str(io.OUT / '独立重建核验.json'), 'used_seconds': budget.used()}, ensure_ascii=False))
    return report


if __name__ == '__main__':
    main()



